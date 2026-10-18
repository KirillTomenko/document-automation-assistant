import asyncio
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from document_automation.extraction import process_paths
from document_automation.generation import (
    DRAFT,
    OperationDenied,
    generate_document,
    verify_document,
)
from document_automation.operations import audit_operation
from document_automation.schemas import ProcessingResult

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("scenario", ["A", "B", "C"])
def test_docx_from_real_demo_files(scenario):
    result = asyncio.run(
        process_paths(tuple((ROOT / "demo" / scenario).glob("*")), scenario)
    )
    result, data = generate_document(result)
    assert verify_document(result, data)
    decision, event = audit_operation(
        "export_document", run_id=scenario, result=result, artifact_bytes=data
    )
    assert decision.allowed and event.event == "operation_decision"
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "{{" not in text
    assert len(doc.tables[0].rows) == 3
    assert doc.tables[0].cell(1, 2).text == "2"
    assert doc.core_properties.identifier == scenario
    assert (DRAFT in text) == (scenario != "A")
    assert (DRAFT in doc.sections[0].header.paragraphs[0].text) == (scenario != "A")
    assert "customer.pdf" not in text and "description.docx" not in text
    assert "items.xlsx" not in text and "format=" not in text
    assert "requested_completion_date" not in text and "unverified_source" not in text
    assert ("Желаемая дата завершения: 20.11.2026" in text) == (scenario == "A")
    assert not audit_operation(
        "export_document", run_id="foreign", result=result, artifact_bytes=data
    )[0].allowed
    assert not verify_document(result, data + b"corrupt")
    # Updating the hash alone cannot hide altered text or missing warning.
    doc.paragraphs[1].text = "APPROVED"
    buffer = BytesIO()
    doc.save(buffer)
    changed = buffer.getvalue()
    raw = result.model_dump()
    raw["generated_document"]["sha256"] = sha256(changed).hexdigest()
    assert not verify_document(ProcessingResult.model_validate(raw), changed)


def test_failed_generation_is_denied(tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"broken")
    result = asyncio.run(process_paths((broken,), "T2"))
    with pytest.raises(OperationDenied) as caught:
        generate_document(result)
    assert not caught.value.event.decision.allowed


@pytest.mark.parametrize("email", ["", "Контактный email: office@example.com"])
def test_no_substantive_content_is_denied(tmp_path, email):
    from business_samples import pdf_snapshot

    from document_automation.contracts import ClassificationResult
    from document_automation.workflows.work_request_inputs import (
        RawCandidate,
        RawFields,
        RegisteredSource,
        WorkRequestInput,
    )
    from document_automation.workflows.work_request_processing import (
        process_work_request,
    )

    parsed = pdf_snapshot("Карточка заказчика\n" + email)
    source = parsed.fragments[0]
    result, _ = process_work_request(
        WorkRequestInput(
            run_id="empty",
            sources=(RegisteredSource(run_id="empty", source_id="0", parsed=parsed),),
            classifications=(
                ClassificationResult(
                    document_id=source.document_id,
                    document_type="customer_card",
                    evidence=(source,),
                ),
            ),
            fields=RawFields(
                contact_email=(
                    RawCandidate(raw_value="office@example.com", source=source),
                )
                if email
                else ()
            ),
        )
    )
    assert result.processing_status == "completed" and result.needs_review
    with pytest.raises(OperationDenied) as caught:
        generate_document(result)
    assert caught.value.event.decision.reason == "no_substantive_content"
