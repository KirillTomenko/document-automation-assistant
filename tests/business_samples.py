"""Synthetic stage-3A inputs, built independently of expected ProcessingResult fixtures."""

from hashlib import sha256

from docx import Document
from openpyxl import Workbook

from document_automation.adapters.parsers import parse_docx, parse_xlsx
from document_automation.contracts import (
    ClassificationResult,
    InputDocument,
    PdfLocation,
    SourceFragment,
)
from document_automation.parsing import ParsedBlock, ParsedDocument
from document_automation.workflows.work_request_inputs import (
    RawCandidate,
    RawFields,
    RawItem,
    RegisteredSource,
    WorkRequestInput,
)


def pdf_snapshot(text):
    """Unicode parser-boundary snapshot; real PDF parsing is covered in test_parsers."""
    digest = sha256(text.encode()).hexdigest()
    doc = InputDocument(
        document_id=f"sha256:{digest}",
        filename="customer.pdf",
        format="pdf",
        sha256=digest,
    )
    location = PdfLocation(page=1)
    fragment = SourceFragment(
        fragment_id=f"{doc.document_id}:pdf/1",
        document_id=doc.document_id,
        text=text,
        location=location,
    )
    return ParsedDocument(
        filename=doc.filename,
        document=doc,
        status="completed",
        fragments=(fragment,),
        blocks=(
            ParsedBlock(
                kind="page",
                location=location,
                text=text,
                fragment_id=fragment.fragment_id,
            ),
        ),
    )


def make_input(
    tmp_path,
    scenario="A",
    *,
    pdf_customer="ООО «Демо Офис»",
    doc_customer="ООО «Демо Офис»",
    doc_date=None,
    excel_date=None,
    email="office@example.com",
    rows=None,
    extra_doc=(),
    mutate_book=None,
    parsed_pdf=None,
):
    pdf = (
        parsed_pdf
        if parsed_pdf is not None
        else pdf_snapshot(
            f"Карточка заказчика\nЗаказчик: {pdf_customer}\nКонтактный email: {email}"
        )
    )
    doc_date = (
        ("" if scenario == "B" else "20.11.2026") if doc_date is None else doc_date
    )
    excel_date = (
        ("" if scenario == "B" else "2026-11-22" if scenario == "C" else "2026-11-20")
        if excel_date is None
        else excel_date
    )
    document = Document()
    for text in (
        "Описание работ",
        "Название заявки: Подготовка офисной комнаты",
        f"Заказчик: {doc_customer}",
        f"Желаемая дата завершения: {doc_date}",
        "Описание работ:",
        "Собрать два офисных стола и установить четыре настенные полки.",
        *extra_doc,
    ):
        document.add_paragraph(text)
    path = tmp_path / "description.docx"
    document.save(path)
    docx = parse_docx(path)
    book = Workbook()
    sheet = book.active
    sheet.title = "Заявка"
    sheet.append(["", ""])
    sheet.append(["Желаемая дата завершения", excel_date])
    sheet = book.create_sheet("Работы")
    sheet.append(["Код позиции", "Наименование работ", "Количество", "Единица"])
    for row in (
        rows
        if rows is not None
        else [
            ["W-01", "Сборка офисного стола", 2, "шт."],
            ["W-02", "Установка настенной полки", 4, "шт"],
        ]
    ):
        sheet.append(row)
    if mutate_book:
        mutate_book(book)
    path = tmp_path / "items.xlsx"
    book.save(path)
    book.close()
    xlsx = parse_xlsx(path)
    fields = {name: [] for name in RawFields.model_fields}
    labels = {
        "customer_name": "Заказчик",
        "contact_email": "Контактный email",
        "request_title": "Название заявки",
        "requested_completion_date": "Желаемая дата завершения",
    }
    for parsed in (pdf, docx):
        for fragment in parsed.fragments:
            for name, label in labels.items():
                for line in fragment.text.splitlines():
                    if line.startswith(label + ":") and (
                        value := line.split(":", 1)[1].strip()
                    ):
                        fields[name].append(
                            RawCandidate(raw_value=value, source=fragment)
                        )
            if fragment.text.startswith("Собрать"):
                fields["work_description"].append(
                    RawCandidate(raw_value=fragment.text, source=fragment)
                )
    for f in xlsx.fragments:
        if (f.location.sheet, f.location.cell) == ("Заявка", "B2"):
            fields["requested_completion_date"].append(
                RawCandidate(raw_value=f.text, source=f)
            )
    items = []
    for sheet in xlsx.sheets:
        if sheet.name == "Работы":
            for row in sheet.table.rows:
                if row.index < 2:
                    continue
                values = {}
                for col, name in zip("ABCD", ("item_code", "name", "quantity", "unit")):
                    fragments = [
                        f
                        for f in xlsx.fragments
                        if f.location.sheet == "Работы"
                        and f.location.cell == f"{col}{row.index}"
                    ]
                    values[name] = tuple(
                        RawCandidate(raw_value=f.text, source=f) for f in fragments
                    )
                # Empty rows are omitted from prepared inputs, as by an extractor.
                if any(values.values()):
                    items.append(
                        RawItem(
                            document_id=xlsx.document.document_id,
                            row=row.index,
                            **values,
                        )
                    )
    parsed_docs = (pdf, docx, xlsx)
    return WorkRequestInput(
        run_id=scenario,
        sources=tuple(
            RegisteredSource(run_id=scenario, source_id=str(i), parsed=p)
            for i, p in enumerate(parsed_docs)
        ),
        classifications=tuple(
            ClassificationResult(
                document_id=p.document.document_id,
                document_type=kind,
                evidence=(
                    tuple(
                        f
                        for f in p.fragments
                        if f.location.format == "xlsx"
                        and f.location.sheet == "Работы"
                        and f.location.cell in ("A1", "B1", "C1", "D1")
                    )
                    if kind == "work_items"
                    else (p.fragments[0],)
                ),
            )
            for p, kind in zip(
                parsed_docs, ("customer_card", "work_description", "work_items")
            )
        ),
        fields=RawFields(**{k: tuple(v) for k, v in fields.items()}),
        items=tuple(items),
    )
