"""Synthetic file builders. No production parser or business extraction is used."""

from datetime import datetime, time, timedelta
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def pdf_file(path, texts=("Demo customer", "Second page"), encrypted=False):
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for text in texts:
        page = writer.add_blank_page(width=300, height=200)
        if text:
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {NameObject("/F1"): writer._add_object(font)}
                    )
                }
            )
            stream = DecodedStreamObject()
            stream.set_data(
                b"BT /F1 12 Tf 20 100 Td (" + text.encode("ascii") + b") Tj ET"
            )
            page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("synthetic-password", algorithm="RC4-128")
    writer.write(path)
    return path


def docx_file(path, nested=True):
    doc = Document()
    doc.add_paragraph("  Синтетический заказчик\tбез нормализации  ")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Код"
    table.cell(0, 1).text = "Описание"
    table.cell(1, 0).text = "W-01"
    table.cell(1, 1).text = "Первая строка\nВторая строка"
    if nested:
        inner = table.cell(1, 1).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "Вложенная таблица"
        table.cell(1, 1).add_paragraph("После вложенной")
    doc.add_paragraph("Конец документа")
    doc.save(path)
    return path


def xlsx_file(path, unsupported=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append(["Код", "Количество", "Дата", "Флаг"])
    ws.append(["  W-01  ", 2, datetime(2026, 11, 20), False])
    ws["B4"] = 2.5
    ws["C4"] = time(12, 30)
    ws["D4"] = timedelta(hours=25)
    ws["A5"] = "00123"
    ws.row_dimensions[4].hidden = True
    ws2 = wb.create_sheet("Пустой")
    ws2.sheet_state = "hidden"
    if unsupported:
        ws["E2"] = "=B2+1"
        ws["F2"] = "Сохранённый якорь"
        ws.merge_cells("F2:G3")
        ws["H2"] = "Независимое значение"
        ws["I2"] = "#DIV/0!"
    wb.save(path)
    wb.close()
    return path


def rewrite_zip(path, transform):
    with ZipFile(path) as source:
        entries = {n: source.read(n) for n in source.namelist()}
    transform(entries)
    with ZipFile(path, "w", ZIP_DEFLATED) as target:
        for name, value in entries.items():
            target.writestr(name, value)


def sample_bytes(builder, tmp_path, name):
    return BytesIO(builder(tmp_path / name).read_bytes())
