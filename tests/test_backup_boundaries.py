"""Filesystem, publication and semantic validation beyond manifest hashes."""

import asyncio
import json
import os
import sqlite3
from hashlib import sha256
from io import BytesIO

import pytest
from docx import Document
from test_backup import (
    clone,
    make_data,
    reseal_database,
    rewrite_manifest,
    tree_bytes,
)
from test_backup import (
    sample_backup as sample_backup,
)
from test_formats_v2 import FIXTURES, batch
from test_llm_usage import adapter_for, tariffs

import document_automation.backup as module
from document_automation.backup_paths import BackupError, checksum
from document_automation.schemas import ProcessingResult
from document_automation.service import DocumentService


@pytest.mark.parametrize("action", ("backup", "restore"))
@pytest.mark.parametrize("fault", ("copy", "publish"))
def test_failure_leaves_target_absent_and_source_unchanged(
    sample_backup, tmp_path, monkeypatch, action, fault
):
    copy, _ = sample_backup
    data = tmp_path / "data"
    module.restore(copy, data)
    source = data if action == "backup" else copy
    before = tree_bytes(source)
    target = tmp_path / "target"
    original = module.copy_file

    def fail_copy(src, dst):
        original(src, dst)
        raise OSError("synthetic partial copy failure")

    def fail_publish(stage, destination):
        raise OSError("synthetic directory publication failure")

    monkeypatch.setattr(
        module,
        "copy_file" if fault == "copy" else "publish",
        fail_copy if fault == "copy" else fail_publish,
    )
    with pytest.raises(OSError):
        (module.backup if action == "backup" else module.restore)(source, target)
    assert not target.exists() and tree_bytes(source) == before
    stages = list(tmp_path.glob(".target-staging-*"))
    assert len(stages) == 1


def test_publication_never_replaces_directory_created_during_restore(
    sample_backup, tmp_path, monkeypatch
):
    source, _ = sample_backup
    target = tmp_path / "target"
    publish = module.publish

    def collision(stage, destination):
        destination.mkdir()
        (destination / "preserve.txt").write_bytes(b"PRESERVE")
        publish(stage, destination)

    monkeypatch.setattr(module, "publish", collision)
    with pytest.raises(BackupError):
        module.restore(source, target)
    assert (target / "preserve.txt").read_bytes() == b"PRESERVE"
    assert not (target / module.DATABASE).exists()


def test_source_database_change_after_snapshot_prevents_backup(tmp_path, monkeypatch):
    data = tmp_path / "data"
    make_data(data, documents=False)
    snapshot = module._sqlite_backup

    def change_source(source, target):
        snapshot(source, target)
        with sqlite3.connect(source / module.DATABASE) as db:
            db.execute("PRAGMA user_version=77")

    monkeypatch.setattr(module, "_sqlite_backup", change_source)
    with pytest.raises(BackupError) as exc:
        module.backup(data, tmp_path / "backup")
    assert exc.value.code == "source_changed" and not (tmp_path / "backup").exists()


def test_directory_junction_cannot_escape_restore_root(sample_backup, tmp_path):
    root, run_id = clone(sample_backup, tmp_path)
    inputs, outside = root / run_id / "inputs", tmp_path / "outside-inputs"
    assert inputs.resolve().is_relative_to(tmp_path.resolve())
    assert outside.absolute().is_relative_to(tmp_path.resolve())
    inputs.rename(outside)
    before = tree_bytes(outside)
    if os.name == "nt":
        import _winapi

        _winapi.CreateJunction(str(outside), str(inputs))
        assert inputs.is_junction()
    else:
        inputs.symlink_to(outside, target_is_directory=True)
    with pytest.raises(BackupError):
        module.verify_backup(root)
    with pytest.raises(BackupError):
        module.restore(root, tmp_path / "restored")
    assert tree_bytes(outside) == before and not (tmp_path / "restored").exists()


def test_resealed_manifest_cannot_register_secret_file(sample_backup, tmp_path):
    root, run_id = clone(sample_backup, tmp_path)
    (root / ".env").write_bytes(b"PRIVATE-KEY")
    size, digest = checksum(root / ".env")
    rewrite_manifest(
        root,
        lambda raw: raw["files"].append(
            dict(path=".env", role="input", run_id=run_id, size=size, sha256=digest)
        ),
    )
    with pytest.raises(BackupError) as exc:
        module.verify_backup(root)
    assert exc.value.code == "manifest_inventory_mismatch"


def test_well_formed_docx_with_resealed_hashes_still_fails_content_verification(
    sample_backup, tmp_path
):
    root, run_id = clone(sample_backup, tmp_path)
    artifact = root / run_id / "work_request.docx"
    doc = Document(artifact)
    doc.paragraphs[0].text = "WRONG DOCUMENT CONTENT"
    buffer = BytesIO()
    doc.save(buffer)
    artifact.write_bytes(buffer.getvalue())
    size, digest = checksum(artifact)
    with sqlite3.connect(root / module.DATABASE) as db:
        raw = json.loads(db.execute("SELECT snapshot FROM runs").fetchone()[0])
        raw["generated_document"]["sha256"] = digest
        result = ProcessingResult.model_validate_json(json.dumps(raw))
        text = result.model_dump_json()
        snapshot_hash = sha256(text.encode()).hexdigest()
        db.execute("UPDATE runs SET snapshot=?,digest=?", (text, snapshot_hash))
        db.execute(
            "UPDATE artifacts SET metadata=?",
            (result.generated_document.model_dump_json(),),
        )
        # A legacy result without artifact_writes is supported, but its bytes
        # must still match the independently rendered document content.
        db.execute("DELETE FROM artifact_writes WHERE kind='document'")

    def mutate(raw):
        next(
            e for e in raw["files"] if e["path"] == f"{run_id}/work_request.docx"
        ).update(size=size, sha256=digest)
        raw["results"][0]["snapshot_sha256"] = snapshot_hash

    rewrite_manifest(root, mutate)
    reseal_database(root)
    assert Document(artifact).core_properties.identifier == run_id
    with pytest.raises(BackupError) as exc:
        module.verify_backup(root)
    assert exc.value.code == "artifact_damaged"


def test_started_attempt_is_not_recovered_or_zeroed_by_backup(tmp_path, monkeypatch):
    adapter, calls = adapter_for()
    service = DocumentService(tmp_path / "data", adapter, tariffs=tariffs())

    def fail_finish(record):
        raise OSError("synthetic ledger finish failure")

    monkeypatch.setattr(service.store, "finish_attempt", fail_finish)
    with pytest.raises(OSError):
        asyncio.run(service.process(batch([FIXTURES / "all.docx"])["files"]))
    before = tree_bytes(service.root)
    with pytest.raises(BackupError) as exc:
        module.backup(service.root, tmp_path / "backup")
    assert exc.value.code == "recovery_required" and tree_bytes(service.root) == before
    assert len(calls) == 1
