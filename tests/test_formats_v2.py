"""Real synthetic files: finite 27-route matrix, evidence, review and both DOCX."""

import asyncio
import base64
import itertools
import json
import shutil
import sqlite3
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from document_automation.adapters.openai_extractor import OpenAIExtractor
from document_automation.adapters.parsers import parse_document
from document_automation.api import create_app
from document_automation.extraction import ExtractionLimits, process_paths
from document_automation.extraction_v2 import DemoExtractorV2
from document_automation.generation import (
    generate_document,
    generate_verification_report,
    verify_document,
    verify_verification_report,
)
from document_automation.schemas import ProcessingResult, result_fingerprint
from document_automation.workflows.work_request_v2_context import inspect_document

FIXTURES = Path(__file__).parent / "fixtures/formats"
ROLES = ("customer", "description", "items")
COMBINATIONS = list(itertools.product(("pdf", "docx", "xlsx"), repeat=3))


def process(paths, extractor=None, limits=None):
    return asyncio.run(
        process_paths(
            tuple(paths), "v2-test", extractor, limits, contract_version="2.0"
        )
    )


def batch(paths):
    return {
        "files": [
            {
                "filename": p.name,
                "content_base64": base64.b64encode(p.read_bytes()).decode(),
            }
            for p in paths
        ]
    }


def expected_values(result):
    assert result.processing_status == "completed"
    assert result.contract_version == "2.0"
    assert not result.needs_review, [(i.code, i.path) for i in result.issues]
    data = result.data
    assert data.customer_name.value == "ООО «Демо Офис»"
    assert data.request_title.value == "Подготовка офисной комнаты"
    assert data.requested_completion_date.value == date(2026, 11, 20)
    assert data.contact_email.value == "office@example.com"
    assert (
        data.work_description.value
        == "Собрать два офисных стола и установить четыре настенные полки."
    )
    assert [
        (i.item_code.value, i.name.value, i.quantity.value, i.unit.value)
        for i in data.items
    ] == [
        ("W-01", "Сборка офисного стола", Decimal("2"), "шт"),
        ("W-02", "Установка настенной полки", Decimal("4"), "шт"),
    ]


def both_docx(result):
    result, work = generate_document(result)
    assert verify_document(result, work)
    doc = Document(BytesIO(work))
    text = "\n".join(p.text for p in doc.paragraphs)
    table = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Подготовка офисной комнаты" in text
    assert "Сборка офисного стола" in table and "W-02" in table
    assert "20.11.2026" in text
    assert len(doc.tables[0].rows) == 3
    result, report = generate_verification_report(result)
    assert verify_verification_report(result, report)
    doc = Document(BytesIO(report))
    text = "\n".join(p.text for p in doc.paragraphs)
    table = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Источники значений и фрагменты" in text
    assert "Сборка офисного стола" in table
    assert "office@example.com" in table
    return result, work, report


@pytest.mark.parametrize("formats", COMBINATIONS, ids=lambda f: "+".join(f))
def test_27_routes_values_provenance_permutations_and_docx(formats):
    paths = tuple(FIXTURES / f"{role}.{fmt}" for role, fmt in zip(ROLES, formats))
    result = process(paths)
    expected_values(result)
    evidence = {
        (f.document_id, f.fragment_id): f
        for p in paths
        for f in inspect_document(parse_document(p)).parsed.fragments
    }
    for field in (result.data.customer_name, result.data.requested_completion_date):
        assert len(field.candidates) == 2
        assert len({c.source.document_id for c in field.candidates}) == 2
    for value in [
        getattr(result.data, n)
        for n in (
            "customer_name",
            "contact_email",
            "request_title",
            "requested_completion_date",
            "work_description",
        )
    ] + [
        getattr(i, n)
        for i in result.data.items
        for n in ("name", "quantity", "unit", "item_code")
    ]:
        for c in value.candidates:
            assert c.verified
            assert evidence[(c.source.document_id, c.source.fragment_id)] == c.source
            assert c.raw_value in c.source.text
    loc = result.data.items[0].quantity.candidates[0].source.location
    if formats[2] == "pdf":
        assert (loc.page, loc.line, loc.column) == (1, 3, 3)
    elif formats[2] == "docx":
        assert (loc.table, loc.row, loc.column) == (1, 2, 3)
    else:
        assert (loc.sheet, loc.cell) == ("Позиции", "C2")
    for permutation in itertools.permutations(paths):
        reordered = process(permutation)
        assert reordered.data == result.data
        assert reordered.classifications == result.classifications
        assert reordered.issues == result.issues
        assert reordered.needs_review == result.needs_review
    both_docx(result)


@pytest.mark.parametrize("fmt", ("pdf", "docx", "xlsx"))
def test_one_file_all_groups(fmt):
    result = process([FIXTURES / f"all.{fmt}"])
    expected_values(result)
    assert {c.document_type for c in result.classifications} == {
        "customer_card",
        "work_description",
        "work_items",
    }
    both_docx(result)


def test_duplicate_file_is_review_without_duplicate_items():
    paths = [FIXTURES / f"{r}.docx" for r in ROLES]
    result = process([*paths, paths[-1]])
    assert len(result.data.items) == 2
    assert result.needs_review
    assert [(i.code, i.path) for i in result.issues] == [
        ("duplicate_document", "documents")
    ]
    assert result.data.customer_name.state == "found"


def test_multiple_item_sources_do_not_merge():
    result = process(
        [
            FIXTURES / "customer.pdf",
            FIXTURES / "description.docx",
            FIXTURES / "items.docx",
            FIXTURES / "items.xlsx",
        ]
    )
    assert len(result.data.items) == 4
    assert result.needs_review
    assert all(i.item_code.state == "ambiguous" for i in result.data.items)
    assert all(i.name.state == "found" for i in result.data.items)
    assert (
        len({i.name.candidates[0].source.document_id for i in result.data.items}) == 2
    )


def test_conflict_partial_and_incomplete(tmp_path):
    path = tmp_path / "items.xlsx"
    shutil.copyfile(FIXTURES / "items.xlsx", path)
    book = load_workbook(path)
    book["Данные"]["B1"] = "2026-11-22"
    book["Позиции"]["C2"] = None
    book.save(path)
    book.close()
    result = process([FIXTURES / "customer.pdf", FIXTURES / "description.docx", path])
    assert result.needs_review
    assert result.data.requested_completion_date.state == "conflict"
    assert result.data.requested_completion_date.value is None
    assert len(result.data.requested_completion_date.candidates) == 2
    assert result.data.items[0].quantity.state == "missing"
    assert result.data.items[0].name.state == "found"
    result, work = generate_document(result)
    assert verify_document(result, work)
    assert result.generated_document.filename == "work_request_needs_review.docx"
    _, report = generate_verification_report(result)
    report_doc = Document(BytesIO(report))
    text = "\n".join(c.text for t in report_doc.tables for r in t.rows for c in r.cells)
    text += "\n".join(p.text for p in report_doc.paragraphs)
    assert "2026-11-22" in text and "20.11.2026" in text
    incomplete = process([FIXTURES / "customer.docx"])
    assert incomplete.needs_review and not incomplete.data.items
    assert incomplete.data.customer_name.state == "found"


@pytest.mark.parametrize("fixture", ("ambiguous.pdf", "unsupported.pdf", "wrapped.pdf"))
def test_unsupported_pdf_is_never_clean_success(fixture):
    result = process(
        [FIXTURES / "customer.xlsx", FIXTURES / "description.docx", FIXTURES / fixture]
    )
    assert result.processing_status == "completed"
    assert result.needs_review
    assert any(
        i.code in ("unsupported_content", "unknown_document") for i in result.issues
    )


@pytest.mark.parametrize("fmt", ("docx", "xlsx"))
def test_ambiguous_and_merged_tables(fmt, tmp_path):
    path = tmp_path / f"items.{fmt}"
    shutil.copyfile(FIXTURES / f"items.{fmt}", path)
    if fmt == "docx":
        doc = Document(path)
        doc.tables[0].cell(1, 0).merge(doc.tables[0].cell(1, 1))
        doc.save(path)
    else:
        book = load_workbook(path)
        book.copy_worksheet(book["Позиции"])
        book.save(path)
        book.close()
    result = process([FIXTURES / "customer.pdf", FIXTURES / "description.docx", path])
    assert result.needs_review
    assert not result.data.items
    assert any(i.code == "unsupported_content" for i in result.issues)


def test_mock_http_v2_and_forged_coordinates():
    calls = []

    async def response(request):
        body = json.loads(request.content)
        assert "row_id" in json.dumps(body["text"]["format"]["schema"])
        calls.append(body)
        output = await DemoExtractorV2().extract(
            body["input"][0]["content"][0]["text"], ExtractionLimits()
        )
        return httpx.Response(
            200,
            json={
                "status": "completed",
                "usage": {"input_tokens": 100, "output_tokens": 50},
                "output": [
                    {
                        "type": "message",
                        "content": [{"type": "output_text", "text": output}],
                    }
                ],
            },
        )

    adapter = OpenAIExtractor(
        "synthetic-test-key",
        "test-model",
        allow_paid=True,
        transport=httpx.MockTransport(response),
        contract_version="2.0",
    )
    result = process([FIXTURES / "all.pdf"], adapter)
    expected_values(result)
    assert len(calls) == 1

    class Forged(DemoExtractorV2):
        async def extract(self, payload, limits):
            output = json.loads(await super().extract(payload, limits))
            output["candidates"].append(
                {
                    "field": "quantity",
                    "document_id": output["candidates"][-1]["document_id"],
                    "fragment_id": output["candidates"][-1]["fragment_id"],
                    "raw_value": "999",
                    "row_id": "invented/row/2",
                }
            )
            return json.dumps(output)

    forged = process([FIXTURES / "all.docx"], Forged())
    assert forged.needs_review
    assert forged.data.items[0].quantity.value == Decimal("2")
    assert any(i.code == "unverified_source" for i in forged.issues)


@pytest.mark.parametrize("formats", COMBINATIONS, ids=lambda f: "+".join(f))
def test_openai_v2_all_27_routes_without_excel_row_rules(formats):
    calls = []

    def response(request):
        body = json.loads(request.content)
        schema = body["text"]["format"]["schema"]["$defs"]["CandidateClaimV2"]
        assert "row_id" in schema["properties"] and "row" not in schema["properties"]
        assert "matching the actual Excel row" not in body["instructions"]
        assert (
            "If the supplied evidence does not establish a valid Excel row"
            not in body["instructions"]
        )
        payload = json.loads(body["input"][0]["content"][0]["text"])
        output = payload["supported_claims"]
        item_claims = [c for c in output["candidates"] if c["row_id"]]
        assert item_claims and all(
            c["row_id"].startswith(formats[2] + "/") for c in item_claims
        )
        calls.append(request)
        return httpx.Response(
            200,
            json={
                "status": "completed",
                "output": [
                    {
                        "type": "message",
                        "content": [
                            {"type": "output_text", "text": json.dumps(output)}
                        ],
                    }
                ],
            },
        )

    adapter = OpenAIExtractor(
        "synthetic-test-key",
        "test-model",
        allow_paid=True,
        transport=httpx.MockTransport(response),
        contract_version="2.0",
    )
    paths = [FIXTURES / f"{role}.{fmt}" for role, fmt in zip(ROLES, formats)]
    result = process(paths, adapter)
    expected_values(result)
    evidence = {
        (f.document_id, f.fragment_id): f
        for p in paths
        for f in inspect_document(parse_document(p)).parsed.fragments
    }
    for field in [
        getattr(result.data, n)
        for n in (
            "customer_name",
            "contact_email",
            "request_title",
            "requested_completion_date",
            "work_description",
        )
    ] + [
        getattr(item, n)
        for item in result.data.items
        for n in ("name", "quantity", "unit", "item_code")
    ]:
        assert field.state == "found"
        for candidate in field.candidates:
            assert candidate.verified
            assert (
                evidence[(candidate.source.document_id, candidate.source.fragment_id)]
                == candidate.source
            )
    assert len(calls) == 1
    both_docx(result)


def test_v1_adapter_rejected_before_v2_network_request():
    calls = []
    adapter = OpenAIExtractor(
        "synthetic-test-key",
        "test-model",
        allow_paid=True,
        transport=httpx.MockTransport(lambda request: calls.append(request)),
    )
    result = process([FIXTURES / "all.docx"], adapter)
    assert result.processing_status == "failed"
    assert result.audit[-1].error_code == "extraction_contract_mismatch"
    assert not calls


def test_restart_reads_v1_v2_without_rewriting_old_snapshot(tmp_path):
    root = tmp_path / "data"
    legacy_paths = [
        Path(__file__).parents[1] / "demo/A" / n
        for n in ("customer.pdf", "description.docx", "items.xlsx")
    ]
    with TestClient(create_app(root, contract_version="1.0")) as client:
        run = client.post("/runs", json=batch(legacy_paths)).json()["run_id"]
        client.post(f"/runs/{run}/document")
        client.post(f"/runs/{run}/verification-report")
        old = client.get(f"/runs/{run}").json()
        work = client.get(f"/runs/{run}/document").content
        report = client.get(f"/runs/{run}/verification-report").content
    with sqlite3.connect(root / "workflow.sqlite3") as db:
        snapshot = db.execute(
            "SELECT snapshot,digest FROM runs WHERE run_id=?", (run,)
        ).fetchone()
    old_fingerprint = result_fingerprint(
        ProcessingResult.model_validate_json(json.dumps(old))
    )
    with TestClient(create_app(root)) as client:
        assert client.get(f"/runs/{run}").json()["data"] == old["data"]
        with sqlite3.connect(root / "workflow.sqlite3") as db:
            assert (
                db.execute(
                    "SELECT snapshot,digest FROM runs WHERE run_id=?", (run,)
                ).fetchone()
                == snapshot
            )
        assert client.get(f"/runs/{run}/document").content == work
        assert client.get(f"/runs/{run}/verification-report").content == report
        legacy = client.get(f"/runs/{run}").json()
        assert "contract_version" not in legacy
        assert legacy["audit"][: len(old["audit"])] == old["audit"]
        assert (
            result_fingerprint(ProcessingResult.model_validate_json(json.dumps(legacy)))
            == old_fingerprint
        )
        new = client.post("/runs", json=batch([FIXTURES / "all.xlsx"])).json()
        new_run = new["run_id"]
        assert new["contract_version"] == "2.0" and not new["needs_review"]
        assert client.post(f"/runs/{new_run}/document").status_code == 200
        assert client.post(f"/runs/{new_run}/verification-report").status_code == 200
    with TestClient(create_app(root)) as client:
        assert client.get(f"/runs/{new_run}/document").status_code == 200
        assert client.get(f"/runs/{new_run}/verification-report").status_code == 200


def test_ten_files_equal_values_and_misleading_names(tmp_path):
    paths = []
    for role, filename in zip(
        ROLES, ("items.docx", "customer.docx", "description.docx")
    ):
        path = tmp_path / filename
        shutil.copyfile(FIXTURES / f"{role}.docx", path)
        paths.append(path)
    for index in range(7):
        doc = Document()
        doc.add_paragraph("Карточка заказчика")
        doc.add_paragraph("Заказчик: ООО «Демо Офис»")
        doc.add_paragraph(f"Синтетическое приложение {index}")
        path = tmp_path / f"attachment-{index}.docx"
        doc.save(path)
        paths.append(path)
    with TestClient(create_app(tmp_path / "store")) as client:
        first = client.post("/runs", json=batch(paths))
        assert first.status_code == 201
        result = ProcessingResult.model_validate_json(first.content)
        expected_values(result)
        assert len(result.documents) == 10
        assert len(result.data.customer_name.candidates) == 9
        reversed_result = client.post("/runs", json=batch(paths[::-1])).json()
        assert reversed_result["data"] == first.json()["data"]
        assert reversed_result["documents"] == first.json()["documents"]
        assert client.post("/runs", json=batch(paths + [paths[0]])).status_code == 422


@pytest.mark.parametrize("fmt", ("docx", "xlsx"))
def test_scalar_pairs_and_table_header_not_description(fmt, tmp_path):
    path = tmp_path / f"pairs.{fmt}"
    pairs = [
        ("Карточка заказчика", ""),
        ("Наименование", "ООО «Демо Офис»"),
        ("Контактный email", "office@example.com"),
        ("Название заявки", "Подготовка офисной комнаты"),
        (
            "Описание работ",
            "Собрать два офисных стола и установить четыре настенные полки.",
        ),
        ("Желаемая дата завершения", "20.11.2026"),
    ]
    if fmt == "docx":
        doc = Document()
        table = doc.add_table(rows=0, cols=2)
        for pair in pairs:
            for cell, text in zip(table.add_row().cells, pair):
                cell.text = text
        doc.save(path)
        item_path = tmp_path / "items.docx"
        doc = Document(FIXTURES / "items.docx")
        doc.tables[0].cell(0, 1).text = "Описание работ"
        doc.tables[0].cell(1, 1).text = "Заказчик: Ложный заказчик внутри позиции"
        doc.save(item_path)
    else:
        from openpyxl import Workbook

        book = Workbook()
        for pair in pairs:
            book.active.append(pair)
        book.save(path)
        book.close()
        item_path = tmp_path / "items.xlsx"
        book = load_workbook(FIXTURES / "items.xlsx")
        book["Позиции"]["B1"] = "Описание работ"
        book["Позиции"]["B2"] = "Заказчик: Ложный заказчик внутри позиции"
        book.save(item_path)
        book.close()
    result = process([path, item_path])
    assert result.data.customer_name.value == "ООО «Демо Офис»"
    assert result.data.work_description.value == pairs[4][1]
    assert len(result.data.work_description.candidates) == 1
    assert len(result.data.items) == 2
    assert not result.needs_review


def test_native_excel_date_and_unsupported_cell(tmp_path):
    path = tmp_path / "all.xlsx"
    book = load_workbook(FIXTURES / "all.xlsx")
    sheet = book["Данные"]
    for row in sheet:
        if row[0].value == "Желаемая дата завершения":
            row[1].value = date(2026, 11, 20)
    book["Позиции"]["C2"] = "=1+1"
    book.save(path)
    book.close()
    result = process([path])
    assert result.needs_review
    assert result.data.requested_completion_date.value == date(2026, 11, 20)
    assert result.data.items[0].quantity.state == "missing"
    assert result.data.items[0].name.state == "found"
    assert any(i.code == "unsupported_content" for i in result.issues)


def test_omitted_conflicting_candidate_fails_without_fallback(tmp_path):
    path = tmp_path / "extra.docx"
    doc = Document()
    doc.add_paragraph("Заказчик: ООО «Другой заказчик»")
    doc.save(path)

    class Omitted(DemoExtractorV2):
        async def extract(self, payload, limits):
            output = json.loads(await super().extract(payload, limits))
            output["candidates"] = [
                c for c in output["candidates"] if "Другой" not in c["raw_value"]
            ]
            return json.dumps(output)

    result = process([FIXTURES / "all.docx", path], Omitted())
    assert result.processing_status == "failed"
    assert result.data is None
    assert result.audit[-1].error_code == "extraction_incomplete_candidates"
    assert not result.can_generate


@pytest.mark.parametrize("failure", ("429", "500", "timeout", "invalid"))
def test_v2_provider_failure_is_safe_no_fallback(failure):
    calls = []

    def response(request):
        calls.append(request)
        if failure == "timeout":
            raise httpx.ReadTimeout("private response detail")
        return httpx.Response(
            200 if failure == "invalid" else int(failure),
            text="private response detail",
        )

    adapter = OpenAIExtractor(
        "synthetic-test-key",
        "test-model",
        allow_paid=True,
        transport=httpx.MockTransport(response),
        contract_version="2.0",
    )
    result = process([FIXTURES / "all.xlsx"], adapter)
    assert result.processing_status == "failed" and result.data is None
    assert len(calls) == 1  # No automatic repeat of a paid or uncertain call.
    assert "private response detail" not in result.model_dump_json()
    assert result.generated_document is None and result.generated_report is None
