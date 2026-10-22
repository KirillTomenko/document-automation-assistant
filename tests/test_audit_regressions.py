"""Stage 10B counterexamples: supported fields and rejected business tables."""

import json
from io import BytesIO

import httpx
import pytest
from docx import Document
from openpyxl import load_workbook
from test_formats_v2 import FIXTURES, expected_values, process

from document_automation.adapters.openai_extractor import OpenAIExtractor
from document_automation.generation import (
    generate_document,
    generate_verification_report,
    verify_document,
    verify_verification_report,
)


def checked_documents(result):
    result, work = generate_document(result)
    assert verify_document(result, work)
    result, report = generate_verification_report(result)
    assert verify_verification_report(result, report)
    assert result.generated_document.filename.endswith(
        "_needs_review.docx" if result.needs_review else "work_request.docx"
    )
    doc = Document(BytesIO(report))
    return "\n".join(
        [p.text for p in doc.paragraphs]
        + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    )


def mock_adapter(calls):
    def response(request):
        calls.append(request)
        body = json.loads(request.content)
        claims = json.loads(body["input"][0]["content"][0]["text"])["supported_claims"]
        return httpx.Response(
            200,
            json={
                "status": "completed",
                "usage": {"input_tokens": 100, "output_tokens": 50},
                "output": [
                    {
                        "type": "message",
                        "content": [
                            {"type": "output_text", "text": json.dumps(claims)}
                        ],
                    }
                ],
            },
        )

    return OpenAIExtractor(
        "synthetic-test-key",
        "test-model",
        allow_paid=True,
        transport=httpx.MockTransport(response),
        contract_version="2.0",
    )


@pytest.mark.parametrize("position", ("before", "after"))
@pytest.mark.parametrize("conflict", (False, True))
@pytest.mark.parametrize("mode", ("offline", "mock"))
def test_labeled_customer_around_xlsx_table(tmp_path, position, conflict, mode):
    book = load_workbook(FIXTURES / "all.xlsx")
    sheet = book["Позиции"]
    raw = "ООО «Другой заказчик»" if conflict else "ООО «Демо Офис»"
    row = 1 if position == "before" else sheet.max_row + 1
    if position == "before":
        sheet.insert_rows(1)
    sheet.cell(row, 1, "Заказчик")
    sheet.cell(row, 2, raw)
    path = tmp_path / "labeled.xlsx"
    book.save(path)
    book.close()
    calls = []
    result = process([path], mock_adapter(calls) if mode == "mock" else None)
    field = result.data.customer_name
    assert result.processing_status == "completed"
    assert result.needs_review == conflict
    assert field.state == ("conflict" if conflict else "found")
    assert field.value == (None if conflict else raw)
    assert len(field.candidates) == 3
    candidate = next(
        c for c in field.candidates if c.source.location.sheet == "Позиции"
    )
    assert candidate.verified and candidate.raw_value == raw
    assert candidate.source.location.cell == f"B{row}"
    assert candidate.source.text == raw
    assert candidate.source.document_id == result.documents[0].document_id
    assert len(result.data.items) == 2
    assert len(calls) == (1 if mode == "mock" else 0)
    report = checked_documents(result)
    assert raw in report and f"B{row}" in report


def test_standalone_inline_label_after_table_and_generic_name_requires_card(tmp_path):
    book = load_workbook(FIXTURES / "all.xlsx")
    sheet = book["Позиции"]
    sheet.append(["Заказчик: ООО «Другой заказчик»"])
    sheet.append(["Наименование", "Не является подтверждённым заказчиком"])
    path = tmp_path / "inline.xlsx"
    book.save(path)
    book.close()
    result = process([path])
    assert result.needs_review and result.data.customer_name.state == "conflict"
    assert len(result.data.customer_name.candidates) == 3
    assert len(result.data.items) == 2
    checked_documents(result)


@pytest.mark.parametrize("extra", (False, True))
@pytest.mark.parametrize("mode", ("offline", "mock"))
def test_partial_xlsx_table_requires_review_even_with_other_items(
    tmp_path, extra, mode
):
    book = load_workbook(FIXTURES / "all.xlsx")
    book["Позиции"]["D1"] = "Неопознанная колонка"
    path = tmp_path / "partial.xlsx"
    book.save(path)
    book.close()
    calls = []
    result = process(
        [path] + ([FIXTURES / "items.docx"] if extra else []),
        mock_adapter(calls) if mode == "mock" else None,
    )
    assert result.processing_status == "completed" and result.needs_review
    issue = next(i for i in result.issues if i.code == "unsupported_content")
    doc = next(d for d in result.documents if d.filename == path.name)
    assert issue.sources and all(
        s.document_id == doc.document_id for s in issue.sources
    )
    assert {s.location.cell for s in issue.sources} == {"A1", "B1", "C1"}
    assert all(s.location.sheet == "Позиции" for s in issue.sources)
    assert len(result.data.items) == (2 if extra else 0)
    assert all(
        i.name.candidates[0].source.document_id != doc.document_id
        for i in result.data.items
    )
    assert len(calls) == (1 if mode == "mock" else 0)
    report = checked_documents(result)
    assert path.name in report and "Позиции" in report


def test_multiple_correct_sources_keep_rows_and_do_not_warn_unsupported():
    result = process([FIXTURES / "all.xlsx", FIXTURES / "items.docx"])
    assert result.needs_review and len(result.data.items) == 4
    assert not any(i.code == "unsupported_content" for i in result.issues)
    assert all(i.item_code.state == "ambiguous" for i in result.data.items)
    checked_documents(result)


def test_auxiliary_table_does_not_trigger_business_table_warning(tmp_path):
    book = load_workbook(FIXTURES / "all.xlsx")
    sheet = book.create_sheet("Справка")
    sheet.append(["Параметр", "Значение", "Единица"])
    sheet.append(["Температура", 20, "C"])
    path = tmp_path / "auxiliary.xlsx"
    book.save(path)
    book.close()
    result = process([path])
    expected_values(result)
    assert not any(i.code == "unsupported_content" for i in result.issues)
    checked_documents(result)


def test_duplicate_header_is_ambiguous_with_source_coordinates(tmp_path):
    book = load_workbook(FIXTURES / "all.xlsx")
    sheet = book["Позиции"]
    sheet["E1"] = "Количество"
    sheet["E2"] = 99
    path = tmp_path / "ambiguous.xlsx"
    book.save(path)
    book.close()
    result = process([path, FIXTURES / "items.docx"])
    assert result.needs_review and len(result.data.items) == 2
    issue = next(i for i in result.issues if i.code == "unsupported_content")
    assert {s.location.cell for s in issue.sources} == {"A1", "B1", "C1", "D1", "E1"}
    checked_documents(result)


def test_item_text_with_scalar_label_does_not_close_table(tmp_path):
    book = load_workbook(FIXTURES / "all.xlsx")
    sheet = book["Позиции"]
    sheet["B2"] = "Заказчик"
    sheet["B3"] = "Заказчик: ООО «Это текст позиции»"
    path = tmp_path / "item-text.xlsx"
    book.save(path)
    book.close()
    result = process([path])
    assert not result.needs_review
    assert result.data.customer_name.value == "ООО «Демо Офис»"
    assert len(result.data.customer_name.candidates) == 2
    assert len(result.data.items) == 2
    checked_documents(result)


def test_partial_table_after_valid_table_does_not_create_fictional_rows(tmp_path):
    book = load_workbook(FIXTURES / "all.xlsx")
    sheet = book["Позиции"]
    sheet.append(["Код работы", "Наименование работ", "Количество", "Неизвестно"])
    sheet.append(["X-99", "Неизвлечённая работа", 9, "шт"])
    path = tmp_path / "two-tables.xlsx"
    book.save(path)
    book.close()
    result = process([path])
    assert result.needs_review and len(result.data.items) == 2
    assert all(i.item_code.value in ("W-01", "W-02") for i in result.data.items)
    issue = next(i for i in result.issues if i.code == "unsupported_content")
    assert {s.location.cell for s in issue.sources} == {"A4", "B4", "C4"}
    checked_documents(result)
