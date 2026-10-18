import base64
import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from document_automation.api import create_app

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("text", ["Unrelated text", "Email: demo@example.com"])
def test_empty_generation_api_and_ui(tmp_path, monkeypatch, text):
    import httpx
    from parsing_samples import pdf_file
    from streamlit.testing.v1 import AppTest

    file = pdf_file(tmp_path / "input.pdf", texts=(text,))
    with TestClient(create_app(tmp_path / "store")) as client:
        result = client.post(
            "/runs",
            json={
                "files": [
                    {
                        "filename": file.name,
                        "content_base64": base64.b64encode(file.read_bytes()).decode(),
                    }
                ]
            },
        ).json()
        run = result["run_id"]
        assert result["processing_status"] == "completed"
        from document_automation.schemas import ProcessingResult

        assert not ProcessingResult.model_validate_json(json.dumps(result)).can_generate
        denial = client.post(f"/runs/{run}/document")
        assert denial.status_code == 409
        assert "содержательных данных" in denial.json()["detail"]
        assert client.get(f"/runs/{run}/document").status_code == 409
        assert client.post(f"/runs/{run}/verification-report").status_code == 200
        assert client.get(f"/runs/{run}/verification-report").status_code == 200
        assert not list((tmp_path / "store" / run).glob("work_request*.docx"))
        assert (
            len(list((tmp_path / "store" / run).glob("verification_report.docx"))) == 1
        )
        audit = client.get(f"/runs/{run}").json()["audit"]
        assert any(
            e.get("decision", {}).get("reason") == "no_substantive_content"
            for e in audit
            if e.get("decision")
        )

        def request(method, url, **kwargs):
            kwargs.pop("timeout", None)
            return client.request(
                method, url.replace("http://127.0.0.1:8000", ""), **kwargs
            )

        monkeypatch.setattr(httpx, "request", request)
        app = AppTest.from_file(
            str(ROOT / "src/document_automation/ui.py"), default_timeout=30
        )
        app.session_state["result"] = result
        app.run()
        assert not app.exception
        assert any("Одного email недостаточно" in e.value for e in app.error)
        assert next(
            b for b in app.button if b.label == "Подготовить рабочую заявку"
        ).disabled
        assert not next(
            b for b in app.button if b.label == "Подготовить отчёт проверки"
        ).disabled


def test_extraction_limits_are_bounded_configuration(tmp_path, monkeypatch):
    from pydantic import ValidationError

    monkeypatch.setenv(
        "DAA_EXTRACTION_LIMITS", '{"attempts":1,"max_output_tokens":2000}'
    )
    app = create_app(tmp_path)
    assert app.state.service.limits.attempts == 1
    assert app.state.service.limits.max_output_tokens == 2000
    monkeypatch.setenv("DAA_EXTRACTION_LIMITS", '{"attempts":999}')
    with pytest.raises(ValidationError):
        create_app(tmp_path)


def payload(scenario="A"):
    return {
        "files": [
            {
                "filename": p.name,
                "content_base64": base64.b64encode(p.read_bytes()).decode(),
            }
            for p in sorted((ROOT / "demo" / scenario).glob("*"))
        ]
    }


@pytest.mark.parametrize("scenario", ["A", "B", "C", "V"])
def test_upload_review_generate_download(tmp_path, scenario):
    with TestClient(create_app(tmp_path)) as client:
        response = client.post("/runs", json=payload(scenario))
        assert response.status_code == 201, response.text
        result = response.json()
        run = result["run_id"]
        assert result["processing_status"] == "completed"
        assert result["needs_review"] == (scenario != "A")
        assert client.get(f"/runs/{run}/document").status_code == 409
        assert client.post(f"/runs/{run}/document").status_code == 200
        download = client.get(f"/runs/{run}/document")
        assert download.status_code == 200
        assert len(Document(BytesIO(download.content)).tables[0].rows) == (
            4 if scenario == "V" else 3
        )
        assert client.post(f"/runs/{run}/verification-report").status_code == 200
        report_response = client.get(f"/runs/{run}/verification-report")
        assert report_response.status_code == 200
        report_doc = Document(BytesIO(report_response.content))
        assert report_doc.core_properties.identifier == run
        report_text = "\n".join(p.text for p in report_doc.paragraphs)
        assert "не устанавливает истинность" in report_text
        assert "Источники значений и фрагменты" in report_text
        assert "create_verification_report: report_allowed" in report_text
        assert "Verification report generated" in report_text
        result_after = client.get(f"/runs/{run}").json()
        assert (
            result_after["generated_document"]["run_id"]
            == result_after["generated_report"]["run_id"]
            == run
        )
        assert client.get(f"/runs/{run}").json()["audit"][-1]["decision"]["allowed"]
        assert client.get("/runs/" + "0" * 32).status_code == 404


@pytest.mark.parametrize(
    "name", ["../evil.pdf", "C:\\evil.pdf", "file.txt", ".hidden.pdf"]
)
def test_reject_upload_paths(tmp_path, name):
    with TestClient(create_app(tmp_path)) as client:
        assert (
            client.post(
                "/runs", json={"files": [{"filename": name, "content_base64": "YQ=="}]}
            ).status_code
            == 400
        )
    assert not list(tmp_path.glob("*/inputs"))


def test_invalid_and_failed_requests(tmp_path):
    with TestClient(create_app(tmp_path)) as client:
        assert client.post("/runs", json={"files": []}).status_code == 422
        assert (
            client.post(
                "/runs",
                json={"files": [{"filename": "x.pdf", "content_base64": "???"}]},
            ).status_code
            == 400
        )
        result = client.post(
            "/runs", json={"files": [{"filename": "x.pdf", "content_base64": "YQ=="}]}
        ).json()
        assert result["processing_status"] == "failed"
        run = result["run_id"]
        assert client.post(f"/runs/{run}/document").status_code == 409
        assert client.get(f"/runs/{run}/document").status_code == 409
        assert client.post(f"/runs/{run}/verification-report").status_code == 409
        assert client.get(f"/runs/{run}/verification-report").status_code == 409


@pytest.mark.parametrize("scenario", ["A", "B", "C", "V"])
def test_streamlit_end_to_end_against_api(tmp_path, monkeypatch, scenario):
    import httpx
    from streamlit.testing.v1 import AppTest

    with TestClient(create_app(tmp_path)) as client:

        def request(method, url, **kwargs):
            kwargs.pop("timeout", None)
            return client.request(
                method, url.replace("http://127.0.0.1:8000", ""), **kwargs
            )

        monkeypatch.setattr(httpx, "request", request)
        app = AppTest.from_file(
            str(ROOT / "src/document_automation/ui.py"), default_timeout=30
        ).run()
        assert not app.exception
        assert "утверждённую структуру" in app.info[0].value
        assert any("OpenAI API ещё не проверено" in c.value for c in app.caption)
        app.selectbox[0].select(scenario).run()
        app.button[0].click().run()
        assert not app.exception
        assert bool(app.warning) == (scenario != "A")
        fields = app.dataframe[0].value
        assert "Позиция 1 — Количество" in fields["Поле"].tolist()
        assert "Найдено" in fields["Состояние"].tolist()
        assert "items[" not in fields.to_string()
        if scenario != "A":
            problems = app.dataframe[1].value
            assert "Что сделать" in problems
            assert "requested_completion_date" not in problems.to_string()
            assert "required_missing" not in problems.to_string()
            assert "conflicting_values" not in problems.to_string()
        assert any("PDF, страница 1" in c.value for c in app.caption)
        assert any(
            (
                "Excel, лист Sheet1, ячейка E3"
                if scenario == "V"
                else "Excel, лист Работы, ячейка C2"
            )
            in c.value
            for c in app.caption
        )
        assert any(e.label == "Технические подробности" for e in app.expander)
        app.button[1].click().run()
        assert not app.exception
        filename, data = app.session_state["document"]
        assert filename == (
            "work_request.docx" if scenario == "A" else "work_request_needs_review.docx"
        )
        assert len(Document(BytesIO(data)).tables[0].rows) == (
            4 if scenario == "V" else 3
        )
        assert len(app.get("download_button")) == 1
        app.button[2].click().run()
        assert not app.exception
        assert {button.label for button in app.get("download_button")} == {
            "Скачать рабочую заявку",
            "Скачать отчёт проверки",
        }
        assert (
            app.session_state["result"]["generated_document"]["run_id"]
            == app.session_state["result"]["generated_report"]["run_id"]
        )
