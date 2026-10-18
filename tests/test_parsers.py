import hashlib
import json
from datetime import datetime, time, timedelta
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import Workbook
from parsing_samples import docx_file, pdf_file, rewrite_zip, xlsx_file
from pydantic import ValidationError
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)

from document_automation.adapters.parsers import (
    parse_document,
    parse_docx,
    parse_pdf,
    parse_xlsx,
)
from document_automation.parsing import ParsedDocument, ParserLimits

BUILDERS = {"pdf": pdf_file, "docx": docx_file, "xlsx": xlsx_file}


def codes(result):
    return {d.code for d in result.diagnostics}


def cell(result, address, sheet=0):
    return next(
        c
        for row in result.sheets[sheet].table.rows
        for c in row.cells
        if c.location.cell == address
    )


@pytest.mark.parametrize("fmt", BUILDERS)
def test_real_files_round_trip_provenance_and_unchanged_bytes(tmp_path, fmt):
    path = BUILDERS[fmt](tmp_path / f"sample.{fmt}")
    original = path.read_bytes()
    result = parse_document(path)
    assert result.status == "completed", result.diagnostics
    assert path.read_bytes() == original
    assert result.document.sha256 == hashlib.sha256(original).hexdigest()
    assert result.fragments
    assert all(
        f.document_id == result.document.document_id and f.location.format == fmt
        for f in result.fragments
    )
    again = parse_document(path)
    assert result == again
    renamed = tmp_path / f"renamed.{fmt}"
    renamed.write_bytes(original)
    assert parse_document(renamed).fragments == result.fragments
    encoded = result.model_dump_json()
    assert ParsedDocument.model_validate_json(encoded) == result
    for mode in ("validation", "serialization"):
        schema = ParsedDocument.model_json_schema(mode=mode)
        Draft202012Validator.check_schema(schema)
        Draft202012Validator(schema, format_checker=FormatChecker()).validate(
            json.loads(encoded)
        )


def test_pdf_pages_and_text(tmp_path):
    result = parse_pdf(pdf_file(tmp_path / "demo.pdf"))
    assert [f.location.page for f in result.fragments] == [1, 2]
    assert [f.text for f in result.fragments] == ["Demo customer", "Second page"]
    assert [b.kind for b in result.blocks] == ["page", "page"]


def test_docx_order_and_nested_tables(tmp_path):
    result = parse_docx(docx_file(tmp_path / "demo.docx"))
    assert [b.kind for b in result.blocks] == ["paragraph", "table", "paragraph"]
    assert result.blocks[0].text == "  Синтетический заказчик\tбез нормализации  "
    table = result.blocks[1].table
    assert table.rows[1].cells[0].blocks[0].text == "W-01"
    nested_cell = table.rows[1].cells[1]
    assert [b.kind for b in nested_cell.blocks] == [
        "paragraph",
        "table",
        "paragraph",
        "paragraph",
    ]
    assert (
        nested_cell.blocks[1].table.rows[0].cells[0].blocks[0].text
        == "Вложенная таблица"
    )
    fragment = next(f for f in result.fragments if f.text == "Вложенная таблица")
    assert (
        fragment.location.table,
        fragment.location.row,
        fragment.location.column,
        fragment.location.paragraph,
    ) == (2, 1, 1, 1)
    assert result.blocks[-1].location.paragraph == 2
    assert result.fragments[-1].text == "Конец документа"


def test_xlsx_values_coordinates_blanks_and_dates(tmp_path):
    result = parse_xlsx(xlsx_file(tmp_path / "demo.xlsx"))
    assert [s.name for s in result.sheets] == ["Данные", "Пустой"]
    assert result.sheets[1].state == "hidden"
    assert cell(result, "A2").value.value == "  W-01  "
    assert cell(result, "A5").value.value == "00123"
    assert cell(result, "B2").value.kind == "integer"
    assert cell(result, "B4").value.value == 2.5
    assert cell(result, "C2").value.value == datetime(2026, 11, 20)
    assert (
        cell(result, "C2").stored_value is not None
    )  # Original Excel serial retained.
    assert cell(result, "C2").number_format
    assert cell(result, "D2").value.value is False
    assert cell(result, "C4").value.value == time(12, 30)
    assert cell(result, "D4").value.value == timedelta(hours=25)
    assert (
        cell(result, "A3").value.kind == "empty"
        and cell(result, "A3").fragment_id is None
    )
    assert [r.index for r in result.sheets[0].table.rows] == [1, 2, 3, 4, 5]
    assert result.sheets[0].table.rows[3].hidden


def test_xlsx_unsupported_does_not_destroy_other_values(tmp_path):
    result = parse_xlsx(xlsx_file(tmp_path / "formula.xlsx", unsupported=True))
    assert result.status == "completed"
    assert "unsupported_content" in codes(result)
    assert cell(result, "E2").value.value == "=B2+1"
    assert not cell(result, "E2").usable and cell(result, "E2").fragment_id is None
    assert cell(result, "F2").value.value == "Сохранённый якорь"
    for address in ("F2", "G2", "F3", "G3"):
        assert (
            not cell(result, address).usable
            and cell(result, address).merged_range == "F2:G3"
        )
    assert cell(result, "I2").value.kind == "error" and not cell(result, "I2").usable
    assert (
        cell(result, "H2").usable
        and cell(result, "H2").value.value == "Независимое значение"
    )


@pytest.mark.parametrize("fmt", BUILDERS)
def test_zero_byte_file_is_failure(tmp_path, fmt):
    path = tmp_path / f"empty.{fmt}"
    path.write_bytes(b"")
    result = parse_document(path)
    assert result.status == "failed" and codes(result) == {"empty_file"}
    assert not result.fragments


@pytest.mark.parametrize("fmt", BUILDERS)
def test_missing_file_is_failure(tmp_path, fmt):
    result = parse_document(tmp_path / f"absent.{fmt}")
    assert result.status == "failed" and codes(result) == {"file_read_error"}


@pytest.mark.parametrize("fmt", BUILDERS)
def test_corrupted_file_is_failure(tmp_path, fmt):
    path = tmp_path / f"corrupt.{fmt}"
    path.write_bytes(b"%PDF-1.7\nbroken" if fmt == "pdf" else b"PK\x03\x04broken")
    result = parse_document(path)
    assert result.status == "failed" and codes(result) == {"file_read_error"}
    assert not result.blocks and not result.sheets and not result.fragments


@pytest.mark.parametrize("fmt", BUILDERS)
def test_file_byte_limit(tmp_path, fmt):
    path = BUILDERS[fmt](tmp_path / f"sample.{fmt}")
    result = parse_document(path, ParserLimits(max_file_bytes=1))
    assert result.status == "failed" and codes(result) == {"resource_limit"}


@pytest.mark.parametrize("suffix", [".txt", ".doc", ".xls", ".png"])
def test_unsupported_extension(tmp_path, suffix):
    result = parse_document(tmp_path / f"file{suffix}")
    assert result.status == "failed" and codes(result) == {"unsupported_format"}


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_mismatched_ooxml_format(tmp_path, fmt):
    other = "xlsx" if fmt == "docx" else "docx"
    path = BUILDERS[other](tmp_path / f"file.{other}")
    target = tmp_path / f"wrong.{fmt}"
    target.write_bytes(path.read_bytes())
    assert codes(parse_document(target)) == {"unsupported_format"}


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_zip_limits_before_library_read(tmp_path, fmt):
    path = BUILDERS[fmt](tmp_path / f"sample.{fmt}")
    for limits in (ParserLimits(max_zip_bytes=1), ParserLimits(max_zip_members=1)):
        assert codes(parse_document(path, limits)) == {"resource_limit"}


def test_encrypted_pdf(tmp_path):
    result = parse_pdf(pdf_file(tmp_path / "encrypted.pdf", encrypted=True))
    assert result.status == "failed" and codes(result) == {"encrypted_file"}


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
@pytest.mark.parametrize("encrypted", [True, False])
def test_office_ole_encryption_marker_detection(tmp_path, fmt, encrypted):
    # Container-signature fixture only: does not test real OOXML decryption.
    path = tmp_path / f"ole.{fmt}"
    data = bytes.fromhex("D0CF11E0A1B11AE1")
    if encrypted:
        data += "EncryptedPackage\0EncryptionInfo".encode("utf-16-le")
    path.write_bytes(data)
    result = parse_document(path)
    assert codes(result) == (
        {"encrypted_file"} if encrypted else {"unsupported_format"}
    )


def test_pdf_blank_and_page_limit(tmp_path):
    path = pdf_file(tmp_path / "blank.pdf", texts=("", ""))
    result = parse_pdf(path)
    assert result.status == "completed" and codes(result) == {
        "no_text",
        "empty_content",
    }
    assert len(result.blocks) == 2
    assert codes(parse_pdf(path, ParserLimits(max_pdf_pages=1))) == {"resource_limit"}


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_valid_empty_document(tmp_path, fmt):
    path = tmp_path / f"empty.{fmt}"
    if fmt == "docx":
        Document().save(path)
    else:
        Workbook().save(path)
    result = parse_document(path)
    assert result.status == "completed" and "empty_content" in codes(result)


def test_xlsx_sheets_cells_and_inflated_range_limits(tmp_path):
    path = xlsx_file(tmp_path / "limits.xlsx")
    assert codes(parse_xlsx(path, ParserLimits(max_xlsx_sheets=1))) == {
        "resource_limit"
    }
    assert codes(parse_xlsx(path, ParserLimits(max_xlsx_cells=2))) == {"resource_limit"}
    wb = Workbook()
    wb.active["XFD1048576"] = "large sparse range"
    wb.save(path)
    assert codes(parse_xlsx(path)) == {"resource_limit"}


def test_docx_nested_depth_limit(tmp_path):
    path = docx_file(tmp_path / "nested.docx")
    assert codes(parse_docx(path, ParserLimits(max_table_depth=1))) == {
        "resource_limit"
    }


@pytest.mark.parametrize("fmt", BUILDERS)
def test_block_and_text_limits(tmp_path, fmt):
    path = BUILDERS[fmt](tmp_path / f"bounded.{fmt}")
    assert codes(parse_document(path, ParserLimits(max_blocks=1))) == {"resource_limit"}
    assert codes(parse_document(path, ParserLimits(max_text_chars=1))) == {
        "resource_limit"
    }


def test_pdf_decoded_stream_limit(tmp_path):
    path = pdf_file(tmp_path / "stream.pdf")
    assert codes(parse_pdf(path, ParserLimits(max_pdf_stream_bytes=1))) == {
        "resource_limit"
    }


def test_docx_merged_text_once(tmp_path):
    path = tmp_path / "merged.docx"
    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(1, 1)).text = "Single anchor"
    doc.save(path)
    result = parse_docx(path)
    assert result.status == "completed" and "unsupported_content" in codes(result)
    assert [f.text for f in result.fragments] == ["Single anchor"]


def test_docx_unsupported_content_warning(tmp_path):
    path = tmp_path / "unsupported.docx"
    doc = Document()
    doc.add_paragraph("Retained text")
    doc.sections[0].header.paragraphs[0].text = "Not extracted header"
    doc.element.body.append(OxmlElement("w:sdt"))
    doc.save(path)
    result = parse_docx(path)
    assert "unsupported_content" in codes(result)
    assert [f.text for f in result.fragments] == ["Retained text"]


@pytest.mark.parametrize(
    "fmt,part", [("docx", "word/document.xml"), ("xlsx", "xl/workbook.xml")]
)
def test_malformed_xml_returns_failure(tmp_path, fmt, part):
    path = BUILDERS[fmt](tmp_path / f"broken.{fmt}")
    rewrite_zip(path, lambda entries: entries.update({part: b"<broken"}))
    assert codes(parse_document(path)) == {"file_read_error"}


def test_schema_rejects_foreign_provenance_and_missing_errors(tmp_path):
    result = parse_pdf(pdf_file(tmp_path / "schema.pdf"))
    data = json.loads(result.model_dump_json())
    data["fragments"][0]["document_id"] = "other"
    with pytest.raises(ValidationError):
        ParsedDocument.model_validate_json(json.dumps(data))
    data = json.loads(result.model_dump_json())
    data["status"] = "failed"
    with pytest.raises(ValidationError):
        ParsedDocument.model_validate_json(json.dumps(data))


def test_changed_file_has_new_identity(tmp_path):
    path = pdf_file(tmp_path / "identity.pdf", texts=("One",))
    before = parse_pdf(path)
    pdf_file(path, texts=("Two",))
    after = parse_pdf(path)
    assert before.document.document_id != after.document.document_id
    assert before.fragments[0].fragment_id != after.fragments[0].fragment_id


def test_pdf_unicode_mapping(tmp_path):
    path = pdf_file(tmp_path / "unicode.pdf", texts=("ABCD",))
    reader = PdfReader(path)
    writer = PdfWriter()
    writer.add_page(reader.pages[0])
    mapping = DecodedStreamObject()
    mapping.set_data(
        b"4 beginbfchar\n<41> <0422>\n<42> <0435>\n<43> <0441>\n<44> <0442>\nendbfchar"
    )
    font = writer.pages[0]["/Resources"]["/Font"]["/F1"].get_object()
    font[NameObject("/ToUnicode")] = writer._add_object(mapping)
    target = tmp_path / "cyrillic.pdf"
    writer.write(target)
    assert parse_pdf(target).fragments[0].text == "Тест"


def test_image_only_pdf_has_explicit_diagnostics(tmp_path):
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    image = DecodedStreamObject()
    image.set_data(b"\xff\x00\x00")
    image.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(1),
            NameObject("/Height"): NumberObject(1),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/XObject"): DictionaryObject(
                {NameObject("/Im0"): writer._add_object(image)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"q 50 0 0 50 0 0 cm /Im0 Do Q")
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "image.pdf"
    writer.write(path)
    result = parse_pdf(path)
    assert result.status == "completed" and not result.fragments
    assert codes(result) == {"unsupported_content", "no_text", "empty_content"}


def test_xlsx_limit_precedes_openpyxl_load(tmp_path, monkeypatch):
    import document_automation.adapters.parsers.xlsx as adapter

    path = xlsx_file(tmp_path / "limits.xlsx")

    def forbidden(*args, **kwargs):
        pytest.fail("openpyxl must not load an over-budget grid")

    monkeypatch.setattr(adapter, "load_workbook", forbidden)
    assert codes(parse_xlsx(path, ParserLimits(max_xlsx_cells=1))) == {"resource_limit"}


def test_bad_cell_address_is_reported(tmp_path):
    path = xlsx_file(tmp_path / "badcell.xlsx")

    def mutate(entries):
        part = "xl/worksheets/sheet1.xml"
        entries[part] = entries[part].replace(b'r="A1"', b'r="BAD"', 1)

    rewrite_zip(path, mutate)
    assert codes(parse_xlsx(path)) == {"file_read_error"}


def test_docx_missing_body_is_failure(tmp_path):
    path = docx_file(tmp_path / "nobody.docx")
    rewrite_zip(
        path,
        lambda entries: entries.update(
            {
                "word/document.xml": b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
            }
        ),
    )
    assert codes(parse_docx(path)) == {"file_read_error"}


def test_xlsx_headers_are_not_silently_ignored(tmp_path):
    path = tmp_path / "header.xlsx"
    wb = Workbook()
    wb.active["A1"] = "Retained"
    wb.active.oddHeader.center.text = "Unsupported header"
    wb.save(path)
    result = parse_xlsx(path)
    assert "unsupported_content" in codes(result)
    assert cell(result, "A1").value.value == "Retained"


@pytest.mark.parametrize("fmt", BUILDERS)
def test_saved_synthetic_examples_match_actual_parser(fmt):
    root = Path(__file__).parent / "fixtures" / "parsing"
    result = parse_document(root / f"sample.{fmt}")
    expected = ParsedDocument.model_validate_json(
        (root / f"{fmt}.parsed.json").read_text(encoding="utf-8")
    )
    assert result == expected
