import asyncio
import json
from io import BytesIO
from pathlib import Path

import httpx
import pytest
from docx import Document
from openpyxl import load_workbook

from document_automation.adapters.openai_extractor import OpenAIExtractor
from document_automation.extraction import (
    DemoExtractor,
    ExtractionLimits,
    process_paths,
)
from document_automation.generation import (
    DRAFT,
    generate_document,
    generate_verification_report,
    verify_document,
)

ROOT = Path(__file__).resolve().parents[1]


def test_variant_real_pipeline_and_docx():
    result = asyncio.run(process_paths(tuple((ROOT / "demo/V").glob("*")), "V"))
    assert result.processing_status == "completed" and result.needs_review
    assert {c.document_type for c in result.classifications} == {
        "customer_card",
        "work_description",
        "work_items",
    }
    data = result.data
    assert data.customer_name.value == "ООО «Демо Офис»"
    assert data.contact_email.value == "office@example.com"
    assert data.customer_name.candidates[0].source.location.format == "pdf"
    assert data.request_title.state == data.requested_completion_date.state == "missing"
    assert (
        data.work_description.value
        == "Собрать учебный стол.\nУстановить учебную полку."
    )
    assert len(data.items) == 3
    assert all(i.item_code.state == "missing" for i in data.items)
    assert data.items[1].quantity.state == "missing"
    assert data.items[1].unit.value == "кв.м"
    assert not any(i.path == "items[1].unit" for i in result.issues)
    assert data.items[0].name.candidates[0].source.location.cell == "D3"
    result, content = generate_document(result)
    assert verify_document(result, content)
    doc = Document(BytesIO(content))
    assert DRAFT in "\n".join(p.text for p in doc.paragraphs)
    assert len(doc.tables[0].rows) == 4
    assert doc.tables[0].cell(1, 0).width.cm >= 2.9
    assert "format=" not in "\n".join(p.text for p in doc.paragraphs)
    assert doc.tables[0].cell(2, 2).text == "кв.м"
    assert "1000" not in data.work_description.value
    assert set(type(data).model_fields) == {
        "customer_name",
        "request_title",
        "work_description",
        "requested_completion_date",
        "contact_email",
        "items",
    }
    result, report = generate_verification_report(result)
    report_doc = Document(BytesIO(report))
    report_text = "\n".join(p.text for p in report_doc.paragraphs)
    assert "Excel, лист Sheet1, ячейка D3" in report_text
    assert "required_missing" in report_text


@pytest.mark.parametrize("order", [(0, 1, 2, 3), (3, 0, 2, 1), (1, 2, 3, 0)])
def test_reordered_columns_and_arbitrary_sheet(tmp_path, order):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Произвольное имя"
    rows = [
        ("Код работы", "Наименование работы", "Кол-во", "Единица измерения"),
        ("W-01", "Сборка стола", 2, "шт"),
    ]
    for row in rows:
        sheet.append([row[i] for i in order])
    path = tmp_path / "items.xlsx"
    book.save(path)
    result = asyncio.run(process_paths((path,), "columns"))
    assert result.data.items[0].name.value == "Сборка стола"
    assert result.data.items[0].item_code.value == "W-01"
    assert result.can_generate


@pytest.mark.parametrize("duplicate", ["header", "sheet"])
def test_ambiguous_table_is_not_selected(tmp_path, duplicate):
    book = load_workbook(ROOT / "demo/V/items.xlsx")
    if duplicate == "header":
        book.active["G2"] = "Количество"
    else:
        book.copy_worksheet(book.active)
    path = tmp_path / "ambiguous.xlsx"
    book.save(path)
    result = asyncio.run(process_paths((path,), "ambiguous"))
    assert not result.data.items and not result.can_generate
    assert any(i.code == "unknown_document" for i in result.issues)


@pytest.mark.parametrize(
    "attack",
    [
        None,
        "object_title",
        "wrong_section",
        "ordinal_code",
        "foreign_id",
        "permissions",
    ],
)
def test_variant_through_mock_openai_adapter(attack):
    async def handler(request):
        body = json.loads(request.content)
        assert body["tools"] == [] and body["text"]["format"]["strict"] is True
        assert "table_headers" in body["instructions"]
        payload = body["input"][0]["content"][0]["text"]
        docs = json.loads(payload)["documents"]
        output = json.loads(await DemoExtractor().extract(payload, ExtractionLimits()))
        if attack == "permissions":
            output["allowed"] = True
        elif attack:
            doc = next(
                d
                for d in docs
                if d["document"]["format"]
                == ("xlsx" if attack == "ordinal_code" else "docx")
            )
            fragment = next(
                f
                for f in doc["fragments"]
                if (
                    f["location"].get("cell") == "A3"
                    if attack == "ordinal_code"
                    else "Наименование объекта" in f["text"]
                )
            )
            output["candidates"].append(
                dict(
                    field="item_code"
                    if attack == "ordinal_code"
                    else "work_description"
                    if attack == "wrong_section"
                    else "request_title",
                    document_id="foreign-run-document"
                    if attack == "foreign_id"
                    else doc["document"]["document_id"],
                    fragment_id=fragment["fragment_id"],
                    raw_value="1" if attack == "ordinal_code" else "Учебный офис",
                    row=3 if attack == "ordinal_code" else None,
                )
            )
        return httpx.Response(
            200,
            json={
                "status": "completed",
                "output": [
                    {
                        "type": "message",
                        "content": [
                            {"type": "output_text", "text": json.dumps(output)}
                        ],
                    }
                ],
            },
        )

    adapter = OpenAIExtractor(
        "synthetic",
        "mock-only",
        allow_paid=True,
        transport=httpx.MockTransport(handler),
    )
    result = asyncio.run(
        process_paths(tuple((ROOT / "demo/V").glob("*")), "mock-V", adapter)
    )
    if attack == "permissions":
        assert result.processing_status == "failed" and not result.can_generate
        return
    assert result.processing_status == "completed"
    assert result.data.request_title.state == "missing"
    assert all(i.item_code.state == "missing" for i in result.data.items)
    assert (
        result.data.work_description.value
        == "Собрать учебный стол.\nУстановить учебную полку."
    )
    if attack:
        assert any(i.code == "unverified_source" for i in result.issues)
    result, content = generate_document(result)
    assert verify_document(result, content)


@pytest.mark.parametrize("label", ["Наименование заказчика", "Организация-заказчик"])
@pytest.mark.parametrize("email_label", ["Email", "E-mail", "Электронная почта"])
def test_pdf_alias_context(label, email_label):
    from business_samples import pdf_snapshot

    from document_automation.workflows.work_request_context import _context_matches
    from document_automation.workflows.work_request_inputs import RawCandidate

    parsed = pdf_snapshot(
        f"{label}: Синтетический заказчик\n{email_label}: demo@example.com"
    )
    for field, value in [
        ("customer_name", "Синтетический заказчик"),
        ("contact_email", "demo@example.com"),
    ]:
        assert _context_matches(
            field, RawCandidate(raw_value=value, source=parsed.fragments[0]), parsed
        )[0]
    assert not _context_matches(
        "request_title",
        RawCandidate(raw_value="Синтетический заказчик", source=parsed.fragments[0]),
        parsed,
    )[0]


@pytest.mark.parametrize(
    "heading,accepted", [("Карточка заказчика", True), ("Счёт", False), ("", False)]
)
def test_generic_name_requires_pdf_customer_card_context(heading, accepted):
    from business_samples import pdf_snapshot

    from document_automation.workflows.work_request_context import _context_matches
    from document_automation.workflows.work_request_inputs import RawCandidate

    parsed = pdf_snapshot(heading + "\nНаименование: Учебная организация")
    candidate = RawCandidate(
        raw_value="Учебная организация", source=parsed.fragments[0]
    )
    assert _context_matches("customer_name", candidate, parsed)[0] is accepted
