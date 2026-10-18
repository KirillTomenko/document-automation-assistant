import asyncio
from collections import Counter
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from business_samples import make_input
from docx import Document
from openpyxl import load_workbook

from document_automation.extraction import process_paths
from document_automation.generation import (
    _clarifications,
    _evidence_registry,
    generate_document,
    generate_verification_report,
    verify_verification_report,
)
from document_automation.workflows.work_request_inputs import RawCandidate
from document_automation.workflows.work_request_processing import process_work_request

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests/fixtures/docx_quality"
ALIASES = [(v, "кв.м") for v in ("кв.м", "кв. м", "кв.м.", "м2", "м²")] + [
    (v, "компл.") for v in ("компл", "компл.", "комплект", "комплекта", "комплектов")
]


def text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


@pytest.mark.parametrize("raw,expected", ALIASES)
def test_unit_alias_through_real_pipeline(tmp_path, raw, expected):
    book = load_workbook(ROOT / "demo/A/items.xlsx")
    book["Работы"]["D2"] = raw
    book["Работы"]["C2"] = 2.5
    path = tmp_path / "items.xlsx"
    book.save(path)
    book.close()
    result = asyncio.run(
        process_paths(
            (ROOT / "demo/A/customer.pdf", ROOT / "demo/A/description.docx", path),
            "alias",
        )
    )
    item = result.data.items[0]
    assert item.unit.value == expected
    assert item.quantity.value == Decimal("2.5")
    candidate = item.unit.candidates[0]
    assert candidate.raw_value == raw and candidate.source.text == raw
    assert candidate.source.location.cell == "D2" and candidate.verified
    assert not result.needs_review
    result, work = generate_document(result)
    assert Document(BytesIO(work)).tables[0].cell(1, 3).text == expected
    result, report = generate_verification_report(result)
    report_text = text(Document(BytesIO(report)))
    assert (
        f"Кандидат: {raw}" in report_text
        and f"После нормализации: {expected}" in report_text
    )
    assert verify_verification_report(result, report)


@pytest.mark.parametrize("unit", ["м³", "кг", "m2", "кв.  м", "часа"])
def test_other_units_remain_invalid(tmp_path, unit):
    result, _ = process_work_request(
        make_input(tmp_path, rows=[["W-01", "Учебная работа", 3, unit]])
    )
    assert ("invalid_value", "items[0].unit") in {
        (i.code, i.path) for i in result.issues
    }
    assert result.needs_review and len(result.data.items) == 1
    result, work = generate_document(result)
    doc = Document(BytesIO(work))
    assert doc.tables[0].cell(1, 3).text == ""
    assert f"«{unit}»" in text(doc) and "допустимы только шт, ч, кв.м, компл." in text(
        doc
    )


def test_missing_unit_is_not_guessed_from_description(tmp_path):
    result, _ = process_work_request(
        make_input(tmp_path, rows=[["W-01", "Покрытие кв.м", 3, None]])
    )
    assert result.data.items[0].unit.state == "missing"
    assert result.data.items[0].quantity.value == 3
    assert result.needs_review


def test_cubic_unit_through_real_pipeline_keeps_partial_row(tmp_path):
    book = load_workbook(FIXTURES / "items.xlsx")
    book.active["B2"] = "м³"
    path = tmp_path / "cubic.xlsx"
    book.save(path)
    book.close()
    result = asyncio.run(
        process_paths(
            (FIXTURES / "customer.pdf", FIXTURES / "description.docx", path), "cubic"
        )
    )
    assert len(result.data.items) == 5 and result.needs_review
    assert result.data.items[0].quantity.value == Decimal("18.5")
    assert result.data.items[0].unit.candidates[0].raw_value == "м³"
    assert ("invalid_value", "items[0].unit") in {
        (i.code, i.path) for i in result.issues
    }
    result, work = generate_document(result)
    doc = Document(BytesIO(work))
    assert len(doc.tables[0].rows) == 6 and doc.tables[0].cell(1, 2).text == ""
    assert "Единица измерения: недопустимое значение «м³»" in text(doc)


def test_invalid_date_stays_manual_and_raw_is_only_a_remark(tmp_path):
    result, _ = process_work_request(
        make_input(tmp_path, doc_date="31.02.2026", excel_date="")
    )
    result, work = generate_document(result)
    doc = Document(BytesIO(work))
    output = text(doc)
    assert "Желаемая дата завершения: __.__.____" in output
    assert "недопустимое значение «31.02.2026»; нужна корректная дата" in output
    assert "Желаемая дата завершения: 31.02.2026" not in output
    assert result.data.requested_completion_date.value is None and result.needs_review


def test_five_positions_placeholders_and_specific_remarks():
    result = asyncio.run(
        process_paths(
            tuple(FIXTURES.glob("*.pdf"))
            + tuple(FIXTURES.glob("*.docx"))
            + tuple(FIXTURES.glob("*.xlsx")),
            "quality",
        )
    )
    assert result.processing_status == "completed" and result.needs_review
    assert len(result.data.items) == 5
    assert (
        result.data.request_title.state
        == result.data.requested_completion_date.state
        == "missing"
    )
    assert [i.unit.value for i in result.data.items] == [
        "кв.м",
        "кв.м",
        "компл.",
        "шт",
        "компл.",
    ]
    assert [i.quantity.value for i in result.data.items] == [
        Decimal("18.5"),
        Decimal(12),
        Decimal(2),
        None,
        Decimal(0),
    ]
    result, work = generate_document(result)
    doc = Document(BytesIO(work))
    assert len(doc.tables[0].rows) == 6
    assert [r.cells[1].text for r in doc.tables[0].rows[1:]] == [
        "18.5",
        "12",
        "2",
        "",
        "",
    ]
    output = text(doc)
    assert "Название заявки: ________________________________" in output
    assert "Желаемая дата завершения: __.__.____" in output
    assert "Сведения о работах" not in output
    assert "items[" not in output and "invalid_value" not in output
    for number, item in enumerate(result.data.items, 1):
        remarks = [
            p.text
            for p in doc.paragraphs
            if p.text.startswith(f"Позиция {number} «{item.name.value}»:")
        ]
        assert len(remarks) == 1 and "Код позиции: не указано" in remarks[0]
    assert "Количество: не указано" in output
    assert "Количество: недопустимое значение «0»; нужно число больше нуля" in output
    # Repeated issues only collapse inside the same position, never across rows.
    repeated = result.model_copy(update={"issues": result.issues + (result.issues[0],)})
    assert _clarifications(repeated) == _clarifications(result)


def test_report_preserves_all_candidates_and_unique_evidence(tmp_path):
    inputs = make_input(tmp_path)
    good = inputs.fields.customer_name[0]
    forged = good.source.model_copy(
        update={"document_id": "foreign", "fragment_id": "claimed"}
    )
    rejected = RawCandidate(raw_value=good.raw_value, source=forged)
    wrong_context = RawCandidate(raw_value="Демо", source=good.source)
    fields = inputs.fields.model_copy(
        update={
            "customer_name": inputs.fields.customer_name
            + (good, rejected, wrong_context)
        }
    )
    result, _ = process_work_request(inputs.model_copy(update={"fields": fields}))
    result, report = generate_verification_report(result)
    doc = Document(BytesIO(report))
    output = text(doc)
    fields = [
        (n, getattr(result.data, n))
        for n in (
            "customer_name",
            "contact_email",
            "request_title",
            "work_description",
            "requested_completion_date",
        )
    ]
    fields += [
        (f"items[{i}].{n}", getattr(item, n))
        for i, item in enumerate(result.data.items)
        for n in type(item).model_fields
    ]
    verified, claimed, reference = _evidence_registry(result, fields)
    assert sum(p.text.startswith("Кандидат:") for p in doc.paragraphs) == sum(
        len(f.candidates) for _, f in fields
    )
    expected = Counter(
        s.text for registry in (verified, claimed) for s, _ in registry.values()
    )
    paragraphs = [p.text for p in doc.paragraphs]
    evidence_texts = paragraphs[
        paragraphs.index("Источники значений и фрагменты") + 1 : paragraphs.index(
            "Что требует участия человека"
        )
    ]
    for fragment, count in expected.items():
        assert evidence_texts.count(fragment) == count
    assert "Неподтверждённые ссылки — не являются доказательствами" in output
    assert reference(forged, False) in output
    assert len(claimed) == 2
    assert reference(good.source) != reference(good.source, False)
    for index in range(1, len(verified) + 1):
        assert sum(p.text == f"Источник {index}" for p in doc.paragraphs) == 1
    assert result.generated_report.audit_snapshot == result.audit
    for event in result.generated_report.audit_snapshot:
        assert event.timestamp.isoformat() in output and event.message in output
    assert verify_verification_report(result, report)


def test_evidence_only_merges_equal_document_location_text(tmp_path):
    inputs = make_input(tmp_path)
    source = inputs.fields.customer_name[0].source
    same = source.model_copy(update={"fragment_id": "same-content-other-id"})
    other_doc = source.model_copy(update={"document_id": "other-document"})
    other_coord = source.model_copy(
        update={"location": source.location.model_copy(update={"page": 2})}
    )
    other_text = source.model_copy(update={"text": source.text + "\nДругой фрагмент"})
    classifications = [
        SimpleNamespace(evidence=(s,))
        for s in (source, same, other_doc, other_coord, other_text)
    ]
    result = SimpleNamespace(classifications=classifications, issues=())
    verified, _, reference = _evidence_registry(result, [])
    assert len(verified) == 4
    assert reference(source) == reference(same)
    assert len(verified[next(iter(verified))][1]) == 2
    assert (
        len({reference(s) for s in (source, other_doc, other_coord, other_text)}) == 4
    )


def test_conflict_does_not_fill_manual_date():
    result = asyncio.run(process_paths(tuple((ROOT / "demo/C").glob("*")), "conflict"))
    result, work = generate_document(result)
    output = text(Document(BytesIO(work)))
    assert "Желаемая дата завершения: __.__.____" in output
    assert "Сверьте источники перед ручным заполнением" in output
    assert "20.11.2026" not in output and "22.11.2026" not in output
    assert (
        result.data.requested_completion_date.state == "conflict"
        and result.needs_review
    )


def test_duplicate_item_code_reason_is_specific_and_preserved(tmp_path):
    result, _ = process_work_request(
        make_input(
            tmp_path,
            rows=[
                ["W-01", "Первая учебная работа", 1, "шт"],
                ["W-01", "Вторая учебная работа", 2, "шт"],
            ],
        )
    )
    result, work = generate_document(result)
    output = text(Document(BytesIO(work)))
    assert output.count("код повторяется у нескольких позиций; уточните коды") == 2
    assert "duplicate_item_code" not in output
    result, report = generate_verification_report(result)
    output = text(Document(BytesIO(report)))
    assert "items[0].item_code: ambiguity_reason=duplicate_item_code" in output
    assert "items[1].item_code: ambiguity_reason=duplicate_item_code" in output
