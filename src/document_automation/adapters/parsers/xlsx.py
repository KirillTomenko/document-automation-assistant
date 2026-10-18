import posixpath
from datetime import date, datetime, time, timedelta
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import coordinate_to_tuple

from document_automation.contracts import XlsxLocation
from document_automation.parsing import (
    BooleanValue,
    DateTimeValue,
    DateValue,
    DurationValue,
    EmptyValue,
    IntegerValue,
    NumberValue,
    ParsedCell,
    ParsedRow,
    ParsedSheet,
    ParsedTable,
    TextValue,
    TimeValue,
)

from .common import ParseFailure, parse_with

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _preflight(data, ctx):
    """Bound the dense grid BEFORE normal-mode openpyxl allocates worksheet objects."""
    stored = {}
    with ZipFile(BytesIO(data)) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        sheets = workbook.findall(f"{NS}sheets/{NS}sheet")
        if not sheets:
            raise ParseFailure("file_read_error", "Workbook contains no sheets.")
        if len(sheets) > ctx.limits.max_xlsx_sheets:
            raise ParseFailure("resource_limit", "XLSX sheet limit exceeded.")
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            r.attrib["Id"]: r.attrib["Target"]
            for r in relations
            if r.attrib.get("TargetMode") != "External"
        }
        total = 0
        for sheet in sheets:
            target = targets[sheet.attrib[REL]]
            target = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath("xl/" + target)
            )
            if not target.startswith("xl/worksheets/"):
                raise ParseFailure(
                    "unsupported_format",
                    "Non-worksheet Excel sheets are not supported.",
                )
            tree = ET.fromstring(archive.read(target))
            if tree.find(f"{NS}headerFooter") is not None:
                ctx.warn(
                    "unsupported_content",
                    f"Sheet {sheet.attrib['name']} has header/footer content; it is not extracted.",
                )
            cells = {}
            max_row = max_col = 1
            for c in tree.findall(f"{NS}sheetData/{NS}row/{NS}c"):
                address = c.attrib["r"]
                if address in cells:
                    raise ParseFailure(
                        "file_read_error", "Duplicate XLSX cell address."
                    )
                row, col = coordinate_to_tuple(address)
                max_row, max_col = max(max_row, row), max(max_col, col)
                cells[address] = (c.findtext(f"{NS}v"), c.attrib.get("t", "n"))
            for merged in tree.findall(f"{NS}mergeCells/{NS}mergeCell"):
                row, col = coordinate_to_tuple(merged.attrib["ref"].split(":")[-1])
                max_row, max_col = max(max_row, row), max(max_col, col)
            total += max_row * max_col
            if total > ctx.limits.max_xlsx_cells:
                raise ParseFailure(
                    "resource_limit", "XLSX total rectangular cell budget exceeded."
                )
            stored[sheet.attrib["name"]] = cells
        extra = [
            n
            for n in archive.namelist()
            if n.startswith(
                (
                    "xl/drawings/",
                    "xl/charts/",
                    "xl/externalLinks/",
                    "xl/comments",
                    "xl/embeddings/",
                )
            )
            or n.endswith("vbaProject.bin")
        ]
        if extra:
            ctx.warn(
                "unsupported_content",
                "Workbook contains drawings, charts, comments, external links or embedded code; these are not extracted or executed.",
            )
    return stored


def _value(cell):
    value = cell.value
    if cell.data_type == "f":
        # Array/data-table formula objects are not ordinary string formula values.
        text = (
            value
            if isinstance(value, str)
            else getattr(value, "text", None) or "[non-scalar formula]"
        )
        return TextValue(kind="formula", value=text)
    if cell.data_type == "e":
        return TextValue(kind="error", value=str(value))
    for typ, model in (
        (bool, BooleanValue),
        (datetime, DateTimeValue),
        (date, DateValue),
        (time, TimeValue),
        (timedelta, DurationValue),
        (int, IntegerValue),
        (float, NumberValue),
    ):
        if isinstance(value, typ):
            return model(value=value)
    if value is None:
        return EmptyValue()
    if isinstance(value, str):
        return TextValue(kind="string", value=value)
    raise ParseFailure("unsupported_format", "Unrecognized Excel cell value type.")


def _read(data, ctx):
    stored = _preflight(data, ctx)
    workbook = load_workbook(
        BytesIO(data), read_only=False, data_only=False, keep_links=False
    )
    sheets = []
    total = 0
    try:
        for sheet in workbook.worksheets:
            total += sheet.max_row * sheet.max_column
            if total > ctx.limits.max_xlsx_cells:
                raise ParseFailure("resource_limit", "XLSX cell budget exceeded.")
            merged = {}
            for region in sheet.merged_cells.ranges:
                ctx.warn(
                    "unsupported_content",
                    f"Merged XLSX range {region} is retained for diagnostics, not extraction.",
                    XlsxLocation(sheet=sheet.title, cell=region.start_cell.coordinate),
                )
                for row in range(region.min_row, region.max_row + 1):
                    for col in range(region.min_col, region.max_col + 1):
                        merged[(row, col)] = str(region)
            rows = []
            for row in sheet.iter_rows():
                cells = []
                for cell in row:
                    ctx.count()
                    location = XlsxLocation(sheet=sheet.title, cell=cell.coordinate)
                    value = (
                        EmptyValue() if isinstance(cell, MergedCell) else _value(cell)
                    )
                    region = merged.get((cell.row, cell.column))
                    usable = region is None and value.kind not in ("formula", "error")
                    if value.kind in ("formula", "error"):
                        ctx.warn(
                            "unsupported_content",
                            f"XLSX {value.kind} is retained without evaluation and is not extraction evidence.",
                            location,
                        )
                    raw, storage_type = stored[sheet.title].get(
                        cell.coordinate, (None, None)
                    )
                    text = (
                        ""
                        if value.value is None
                        else value.value.isoformat()
                        if isinstance(value.value, (date, datetime, time))
                        else str(value.value)
                    )
                    fid = (
                        ctx.fragment(
                            text, location, f"xlsx/{sheet.title}/{cell.coordinate}"
                        )
                        if usable
                        else None
                    )
                    cells.append(
                        ParsedCell(
                            location=location,
                            value=value,
                            number_format=cell.number_format,
                            stored_value=raw,
                            storage_type=storage_type,
                            fragment_id=fid,
                            usable=usable,
                            merged_range=region,
                        )
                    )
                rows.append(
                    ParsedRow(
                        index=row[0].row,
                        cells=tuple(cells),
                        hidden=bool(sheet.row_dimensions[row[0].row].hidden),
                    )
                )
            sheets.append(
                ParsedSheet(
                    name=sheet.title,
                    state=sheet.sheet_state,
                    table=ParsedTable(table_id=f"xlsx/{sheet.title}", rows=tuple(rows)),
                )
            )
    finally:
        workbook.close()
    return [], sheets


def parse_xlsx(path, limits=None):
    return parse_with(path, "xlsx", _read, limits)
