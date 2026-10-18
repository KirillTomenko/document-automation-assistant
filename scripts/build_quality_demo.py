"""Invented five-position equivalent; never reads the private customer documents."""

import os
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]


def build(directory):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    directory.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(
        TTFont(
            "Synthetic",
            os.environ.get("DAA_FIXTURE_FONT", "C:/Windows/Fonts/arial.ttf"),
        )
    )
    pdf = canvas.Canvas(str(directory / "customer.pdf"))
    pdf.setAuthor("Invented regression fixture")
    pdf.setFont("Synthetic", 12)
    for index, line in enumerate(
        (
            "Карточка заказчика",
            "Заказчик: ООО «СтройМастер Демо»",
            "Email: training@example.com",
        )
    ):
        pdf.drawString(45, 780 - index * 24, line)
    pdf.save()
    doc = Document()
    doc.core_properties.author = "Invented regression fixture"
    doc.add_paragraph("Наименование объекта: Учебный стенд")
    doc.add_paragraph(
        "Этапы работ:\nПодготовить учебное основание и уложить покрытие.\nСобрать и установить учебные элементы.\nСтоимость: не извлекать"
    )
    doc.save(directory / "description.docx")
    book = Workbook()
    sheet = book.active
    sheet.title = "Учебный перечень"
    for row in (
        ("№", "Ед. изм.", "Наименование работ", "Кол-во"),
        (1, "кв. м", "Подготовка учебного основания", 18.5),
        (2, "м²", "Укладка учебного покрытия", 12),
        (3, "комплектов", "Сборка учебного монтажного комплекта", 2),
        (4, "шт.", "Установка учебной панели", None),
        (5, "компл", "Монтаж учебного узла", 0),
    ):
        sheet.append(row)
    book.properties.creator = "Invented regression fixture"
    book.save(directory / "items.xlsx")
    book.close()


if __name__ == "__main__":
    build(ROOT / "tests/fixtures/docx_quality")
