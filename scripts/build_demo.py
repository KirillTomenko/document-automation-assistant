"""Rebuild only synthetic demo inputs and the fixed application template."""

from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]


def main():
    assets = ROOT / "src/document_automation/assets"
    assets.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(22)
    doc.add_paragraph("Сводная заявка на выполнение работ", "Title")
    doc.add_paragraph("{{status}}")
    doc.add_paragraph(
        "Документ сформирован автоматически и подлежит проверке человеком"
    )
    for title, key in [
        ("Название заявки", "request_title"),
        ("Заказчик", "customer_name"),
        ("Контактный email", "contact_email"),
        ("Желаемая дата завершения", "requested_completion_date"),
    ]:
        doc.add_paragraph(title + ": {{" + key + "}}")
    doc.add_heading("Описание работ", 1)
    doc.add_paragraph("{{work_description}}")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text, width in zip(
        table.rows[0].cells,
        ("Код", "Наименование работ", "Количество", "Единица"),
        (2, 10, 3, 2.4),
    ):
        cell.text = text
        cell.width = Cm(width)
    table.autofit = False
    doc.add_heading("Проблемы, требующие проверки", 1)
    doc.add_paragraph("{{issues}}")
    doc.add_page_break()
    doc.add_heading("Источники данных", 1)
    doc.add_paragraph("{{sources}}")
    section.header.paragraphs[0].text = "{{header}}"
    doc.core_properties.author = "Document Automation Assistant"
    doc.save(assets / "work_request.docx")
    for scenario in ("A", "B", "C"):
        directory = ROOT / "demo" / scenario
        directory.mkdir(parents=True, exist_ok=True)
        copyfile(
            ROOT / "tests/fixtures/business/customer.pdf", directory / "customer.pdf"
        )
        description = Document()
        for text in (
            "Описание работ",
            "Название заявки: Подготовка офисной комнаты",
            "Заказчик: ООО «Демо Офис»",
            "Желаемая дата завершения: " + ("" if scenario == "B" else "20.11.2026"),
            "Описание работ:",
            "Собрать два офисных стола и установить четыре настенные полки.",
        ):
            description.add_paragraph(text)
        description.core_properties.author = "Synthetic demo"
        description.save(directory / "description.docx")
        book = Workbook()
        book.active.title = "Заявка"
        book.active["A2"] = "Желаемая дата завершения"
        book.active["B2"] = (
            "" if scenario == "B" else "2026-11-22" if scenario == "C" else "2026-11-20"
        )
        sheet = book.create_sheet("Работы")
        for row in [
            ("Код позиции", "Наименование работ", "Количество", "Единица"),
            ("W-01", "Сборка офисного стола", 2, "шт."),
            ("W-02", "Установка настенной полки", 4, "шт"),
        ]:
            sheet.append(row)
        book.properties.creator = "Synthetic demo"
        book.save(directory / "items.xlsx")
        book.close()


if __name__ == "__main__":
    main()
