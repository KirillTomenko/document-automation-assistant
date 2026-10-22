"""Build invented v2 fixtures; optional Matplotlib is needed only to rebuild PDFs."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ["Код позиции", "Наименование работ", "Количество", "Единица"]
ROWS = [
    ["W-01", "Сборка офисного стола", 2, "шт."],
    ["W-02", "Установка настенной полки", 4, "шт"],
]
CUSTOMER = [
    "Карточка заказчика",
    "Заказчик: ООО «Демо Офис»",
    "Контактный email: office@example.com",
]
DESCRIPTION = [
    "Описание работ",
    "Название заявки: Подготовка офисной комнаты",
    "Заказчик: ООО «Демо Офис»",
    "Желаемая дата завершения: 20.11.2026",
    "Описание работ:",
    "Собрать два офисных стола и установить четыре настенные полки.",
]
ITEMS = [
    "Желаемая дата завершения: 2026-11-20",
    " | ".join(HEADERS),
    *(" | ".join(map(str, row)) for row in ROWS),
]


def pdf(path, lines):
    from matplotlib import rc_context
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.figure import Figure

    with rc_context({"pdf.fonttype": 42}):
        figure = Figure(figsize=(8.27, 11.69))
        for index, line in enumerate(lines):
            figure.text(
                0.05, 0.95 - index * 0.028, line, fontsize=9, fontfamily="DejaVu Sans"
            )
        with PdfPages(
            path,
            metadata={
                "Title": "Synthetic v2 fixture",
                "Author": "Synthetic regression fixture",
                "CreationDate": None,
                "ModDate": None,
            },
        ) as output:
            output.savefig(figure)
    # Matplotlib emits empty indirect annotation/XObject containers. The runtime
    # deliberately warns on such entries; omit only containers verified empty.
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(path)
    writer = PdfWriter()
    for page in reader.pages:
        if "/Annots" in page and not page["/Annots"]:
            del page["/Annots"]
        resources = page["/Resources"]
        if "/XObject" in resources and not resources["/XObject"]:
            del resources["/XObject"]
        writer.add_page(page)
    writer.add_metadata(
        {"/Title": "Synthetic v2 fixture", "/Author": "Synthetic regression fixture"}
    )
    writer.write(path)


def build(directory):
    directory.mkdir(parents=True, exist_ok=True)
    for role, lines in (
        ("customer", CUSTOMER),
        ("description", DESCRIPTION),
        ("items", ITEMS),
        ("all", CUSTOMER + DESCRIPTION + ITEMS),
    ):
        pdf(directory / f"{role}.pdf", lines)
        doc = Document()
        doc.core_properties.author = "Synthetic regression fixture"
        for line in (
            CUSTOMER
            if role == "customer"
            else DESCRIPTION
            if role == "description"
            else [ITEMS[0]]
            if role == "items"
            else CUSTOMER + DESCRIPTION + [ITEMS[0]]
        ):
            doc.add_paragraph(line)
        if role in ("items", "all"):
            table = doc.add_table(rows=1, cols=4)
            for c, text in zip(table.rows[0].cells, HEADERS):
                c.text = text
            for values in ROWS:
                for c, text in zip(table.add_row().cells, values):
                    c.text = str(text)
        doc.save(directory / f"{role}.docx")
        book = Workbook()
        book.properties.creator = "Synthetic regression fixture"
        sheet = book.active
        sheet.title = "Данные"
        for line in (
            CUSTOMER
            if role == "customer"
            else DESCRIPTION
            if role == "description"
            else [ITEMS[0]]
            if role == "items"
            else CUSTOMER + DESCRIPTION + [ITEMS[0]]
        ):
            if ":" in line:
                label, value = line.split(":", 1)
                if label == "Описание работ" and not value.strip():
                    continue
                sheet.append([label, value.strip() or None])
            elif line.startswith("Собрать"):
                sheet.append(["Описание работ", line])
            else:
                sheet.append([line])
        if role in ("items", "all"):
            sheet = book.create_sheet("Позиции")
            sheet.append(HEADERS)
            for row in ROWS:
                sheet.append(row)
        book.save(directory / f"{role}.xlsx")
        book.close()
    pdf(directory / "ambiguous.pdf", ITEMS + ["", *ITEMS[1:]])
    pdf(
        directory / "unsupported.pdf",
        ["Перечень работ", "Наименование работ  Количество  Единица", "Работа  2  шт"],
    )
    pdf(
        directory / "wrapped.pdf",
        [ITEMS[1], ITEMS[2], "Продолжение строки без разделителя", ITEMS[3]],
    )


if __name__ == "__main__":
    build(ROOT / "tests/fixtures/formats")
